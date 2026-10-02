from pathlib import Path
import json, re, math, hashlib, unicodedata
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WORK=Path(__file__).resolve().parent
OUT=WORK.parent/'MeBOD_115年系統手冊_物件導向詳細版.docx'
FIG=WORK/'figures'; FIG.mkdir(exist_ok=True)
FONT='C:/Windows/Fonts/kaiu.ttf'
F=lambda n: ImageFont.truetype(FONT,n)
W=1800

def canvas(h=1200):
    im=Image.new('RGB',(W,h),'white'); return im,ImageDraw.Draw(im)

def wrap(s,max_width,font):
    out=[]
    for line in s.split('\n'):
        cur=''
        for ch in line:
            if font.getlength(cur+ch)>max_width and cur:
                out.append(cur);cur=ch
            else: cur+=ch
        out.append(cur)
    return out

def txt(d,xy,text,size=34,width=400,center=True):
    lines=wrap(text,width,F(size)); x,y=xy
    for i,line in enumerate(lines):
        d.text((x,y+i*(size+10)),line,font=F(size),fill='black',anchor='mt' if center else 'lt')

def box(d,rect,text,kind='box',size=34):
    x,y,w,h=rect
    if kind=='ellipse': d.ellipse((x,y,x+w,y+h),fill='#f5f5f5',outline='black',width=3)
    elif kind=='diamond': d.polygon([(x+w/2,y),(x+w,y+h/2),(x+w/2,y+h),(x,y+h/2)],fill='#f5f5f5',outline='black',width=3)
    else: d.rounded_rectangle((x,y,x+w,y+h),radius=14 if kind=='state' else 0,fill='#f5f5f5',outline='black',width=3)
    inset=0
    if kind=='box' and w>=400:
        icon_kind=next((k for words,k in [(['PostgreSQL','資料庫','SQLAlchemy'],'database'),(['Gemini','AI','模型'],'ai'),(['瀏覽器','Vue','Dashboard','LoginView'],'browser'),(['爬蟲','來源','平台'],'web'),(['使用者','帳號'],'user')] if any(s in text for s in words)), 'server')
        pictogram(d,x+18,y+(h-62)/2,icon_kind)
        inset=80
    lines=wrap(text,w-32-inset,F(size)); yy=y+(h-len(lines)*(size+10))/2
    txt(d,(x+inset+(w-inset)/2,yy),text,size,w-32-inset)

def pictogram(d,x,y,kind):
    c='#315C80'; pale='#EAF2F8'
    if kind=='database':
        d.rectangle((x+5,y+13,x+57,y+49),fill=pale,outline=c,width=3)
        d.ellipse((x+5,y+37,x+57,y+60),fill=pale,outline=c,width=3)
        d.rectangle((x+8,y+28,x+54,y+46),fill=pale)
        d.ellipse((x+5,y+2,x+57,y+25),fill=pale,outline=c,width=3)
        d.arc((x+5,y+22,x+57,y+45),0,180,fill=c,width=3)
    elif kind=='user':
        d.ellipse((x+20,y+2,x+44,y+26),fill=pale,outline=c,width=3)
        d.arc((x+5,y+28,x+59,y+81),180,360,fill=c,width=4)
        d.line((x+5,y+55,x+59,y+55),fill=c,width=3)
    elif kind=='web':
        d.ellipse((x+3,y+3,x+59,y+59),outline=c,width=3)
        d.ellipse((x+18,y+3,x+44,y+59),outline=c,width=2)
        d.line((x+4,y+31,x+59,y+31),fill=c,width=2)
    elif kind=='ai':
        d.rounded_rectangle((x+10,y+10,x+52,y+52),radius=6,fill=pale,outline=c,width=3)
        for n in [19,31,43]:
            d.line((x+n,y+2,x+n,y+10),fill=c,width=3);d.line((x+n,y+52,x+n,y+60),fill=c,width=3)
            d.line((x+2,y+n,x+10,y+n),fill=c,width=3);d.line((x+52,y+n,x+60,y+n),fill=c,width=3)
        d.text((x+19,y+20),'AI',font=F(22),fill=c)
    elif kind=='browser':
        d.rounded_rectangle((x+2,y+5,x+60,y+50),radius=4,fill=pale,outline=c,width=3)
        d.line((x+3,y+17,x+59,y+17),fill=c,width=2)
        for n in [9,17,25]:d.ellipse((x+n,y+10,x+n+2,y+12),fill=c)
        d.line((x+31,y+50,x+31,y+59),fill=c,width=3)
        d.line((x+16,y+59,x+46,y+59),fill=c,width=3)
    else:
        for n in range(3):
            d.rounded_rectangle((x+4,y+3+n*20,x+58,y+18+n*20),radius=3,fill=pale,outline=c,width=2)
            d.ellipse((x+11,y+8+n*20,x+16,y+13+n*20),fill=c)

def arrow(d,start,end,label='',dashed=False):
    x,y=start;xx,yy=end
    if dashed:
        length=math.hypot(xx-x,yy-y)
        for n in range(0,int(length),18):
            a=n/length;b=min((n+10)/length,1)
            d.line((x+(xx-x)*a,y+(yy-y)*a,x+(xx-x)*b,y+(yy-y)*b),fill='black',width=3)
    else:d.line((start,end),fill='black',width=3)
    angle=math.atan2(yy-y,xx-x)
    d.polygon([(xx,yy),(xx-18*math.cos(angle-.5),yy-18*math.sin(angle-.5)),(xx-18*math.cos(angle+.5),yy-18*math.sin(angle+.5))],fill='black')
    if label:
        mid=((x+xx)/2,(y+yy)/2-37)
        font=F(29); tw=font.getlength(label)
        d.rectangle((mid[0]-tw/2-6,mid[1]-3,mid[0]+tw/2+6,mid[1]+32),fill='white')
        txt(d,mid,label,29,900)

def flow(name,nodes,edges,h=1200):
    im,d=canvas(h)
    for a,b,label in edges:
        ra=nodes[a][0];rb=nodes[b][0]
        ax,ay,aw,ah=ra;bx,by,bw,bh=rb
        if abs((ax+aw/2)-(bx+bw/2))>abs((ay+ah/2)-(by+bh/2)):
            st=(ax+aw if bx>ax else ax,ay+ah/2);en=(bx if bx>ax else bx+bw,by+bh/2)
        else:st=(ax+aw/2,ay+ah if by>ay else ay);en=(bx+bw/2,by if by>ay else by+bh)
        if any(e[0]==b and e[1]==a for e in edges):
            offset=-40 if a<b else 40
            if st[1]==en[1]:st=(st[0],st[1]+offset);en=(en[0],en[1]+offset)
            else:st=(st[0]+offset,st[1]);en=(en[0]+offset,en[1])
        arrow(d,st,en,label)
    for rect,text,*kind in nodes.values():box(d,rect,text,kind[0] if kind else 'box')
    im.save(FIG/(name+'.png'))

def seq(name,actors,events):
    h=220+len(events)*125;im,d=canvas(h); xs=[140+i*(1520/(len(actors)-1)) for i in range(len(actors))]
    for x,t in zip(xs,actors):
        box(d,(x-130,25,260,110),t,size=30)
        d.line((x,135,x,h-30),fill='#777',width=2)
    for i,(a,b,t) in enumerate(events):
        y=225+i*125
        if a==b:
            d.line((xs[a],y,xs[a]+70,y,xs[a]+70,y+45,xs[a],y+45),fill='black',width=3)
            txt(d,(xs[a]+90,y-15),t,28,360,False)
        else:arrow(d,(xs[a],y),(xs[b],y),t,dashed=b<a)
    im.save(FIG/(name+'.png'))

flow('architecture',{
'u':((620,30,560,140),'瀏覽器\nVue 單頁應用'),
'a':((620,300,560,150),'FastAPI 路由與授權'),
's':((620,570,560,150),'查詢 分析 問答 匯出服務'),
'p':((60,880,470,160),'PostgreSQL\n資料與租約'),
'g':((660,880,470,160),'Gemini\n文字與語意向量'),
'c':((1240,570,490,150),'排程與四平台爬蟲'),
'e':((1240,880,490,160),'PTT Dcard\nMobile01 Threads')},
[('u','a','HTTPS / JSON'),('a','s','呼叫'),('s','p','SQL'),('s','g','API'),('a','c','管理員啟動'),('c','e','取得資料')],1100)
flow('pipeline',{str(i):((250+([0,1,1,0,0,1][i])*950,60+(i//2)*310,550,170),t) for i,t in enumerate(['來源取得','相關性過濾','文章更新與留言同步','情緒判讀','段落向量產生','監測與預警'])},[('0','1',''),('1','2',''),('2','3',''),('3','4',''),('4','5','')],1000)
# Actor/use-case diagram: boundaries and associations, not data flow.
im,d=canvas(1260);d.rectangle((470,30,1740,1200),outline='black',width=3);txt(d,(1100,50),'MeBOD 系統',40,1000)
actors=[('訪客',140),('一般帳號',530),('管理員',950)]
for name,y in actors:
    d.ellipse((190,y,250,y+60),outline='black',width=3);d.line((220,y+60,220,y+150),fill='black',width=3);d.line((150,y+95,290,y+95),fill='black',width=3);d.line((220,y+150,165,y+210),fill='black',width=3);d.line((220,y+150,275,y+210),fill='black',width=3);txt(d,(220,y+230),name,36,330)
cases=[(600,135,'瀏覽儀表板與文章'),(1160,135,'登入與登出'),(600,460,'AI問答與Excel'),(1160,460,'個人歷史與監測'),(600,850,'帳號與系統設定'),(1160,850,'爬取與維護')]
for x,y,t in cases:box(d,(x,y,500,170),t,'ellipse')
for st,en in [((290,235),(600,220)),((290,625),(600,220)),((290,625),(600,545)),((290,625),(1160,545)),((290,1045),(600,935)),((290,1045),(1160,935))]:d.line((st,en),fill='#333',width=2)
txt(d,(1100,1110),'一般帳號功能仍依方案限制；管理員可維運',30,1100)
im.save(FIG/'usecase.png')
flow('activity_qa',{'a':((620,20,550,120),'輸入問題與脈絡','state'),'b':((620,245,550,170),'登入及額度有效','diamond'),'x':((40,260,410,140),'回401或403\n顯示原因','state'),'c':((620,545,550,130),'快取或混合檢索'),'d':((620,780,550,130),'模型生成及結構驗證'),'e':((620,1015,550,130),'成功回應並記錄用量','state')},[('a','b',''),('b','x','否'),('b','c','是'),('c','d',''),('d','e','')],1200)
flow('activity_crawl',{'a':((640,20,530,130),'排程或管理員啟動','state'),'b':((640,240,530,160),'取得租約','diamond'),'x':((50,250,430,140),'已有工作\n拒絕重複執行','state'),'c':((640,520,530,130),'爬取及相關性過濾'),'d':((640,760,530,160),'續租仍有效','diamond'),'y':((1230,760,520,160),'失去租約\n停止與丟棄批次','state'),'e':((640,1050,530,140),'更新 分析 記錄 釋放','state')},[('a','b',''),('b','x','否'),('b','c','是'),('c','d',''),('d','y','否'),('d','e','是')],1240)
flow('analysis_classes',{'b':((60,100,480,400),'«boundary»\nLoginView\nDashboardView\nHistoryView\nQAView'),'c':((660,100,480,400),'«control»\n身分與權限\n查詢與分析\n問答與監測'),'e':((1260,100,480,400),'«entity»\nUser Article\nAnalysisHistory\nWatchKeyword Alert')},[('b','c','請求'),('c','e','存取')],630)
seq('seq_login',['使用者','前端與路由','Auth API','PostgreSQL'],[(0,1,'輸入帳號與密碼'),(1,2,'POST login'),(2,3,'讀取帳號及驗證狀態'),(3,2,'帳號與密碼雜湊'),(2,1,'Set-Cookie 與使用者資訊'),(1,0,'進入 Dashboard'),(0,1,'按下登出'),(1,2,'等待 POST logout'),(2,1,'清除 Cookie'),(1,0,'清除本機狀態後導向登入')])
seq('seq_qa',['使用者','QA API','RAG 服務','資料庫','Gemini'],[(0,1,'問題與上下文'),(1,2,'驗證登入 額度 與限流'),(2,3,'查詢快取與候選文章'),(2,4,'問題向量與意圖'),(4,2,'向量與查詢條件'),(2,3,'取得文章段落向量'),(2,4,'融合排名後的來源與問題'),(4,2,'結構化回答'),(2,3,'快取結果'),(1,3,'成功用量計次'),(1,0,'回答與來源')])
seq('seq_crawl',['排程器','租約服務','爬蟲','文章服務','資料庫'],[(0,1,'取得 crawler 租約'),(1,4,'原子寫入或過期接手'),(0,2,'啟用平台及看板'),(2,0,'正規化文章批次'),(0,1,'續租與確認owner'),(0,3,'過濾後新增或更新'),(3,4,'文章 留言 失效舊分析'),(0,4,'情緒 向量 預警 日誌'),(0,1,'依owner釋放')])
flow('design_classes',{'r':((60,80,470,330),'«router»\nanalysis qa export\n驗證HTTP輸入\n注入Session與User'),'s':((660,80,480,330),'«service»\nrag dashboard plan\n篩選 計算 授權\n處理外部失敗'),'m':((1260,80,480,330),'«model»\nSQLAlchemy\nArticle User History\n主鍵 外鍵 約束'),'l':((660,620,480,250),'«adapter»\nllm_client\nGemini回應處理')},[('r','s','呼叫'),('s','m','查詢與更新'),('s','l','生成')],960)
flow('deployment',{'u':((90,80,600,190),'«device» 使用者電腦\n瀏覽器'),'p':((1050,80,650,190),'«node» HTTPS代理\nTLS及可信來源位址'),'a':((1050,510,650,210),'«executionEnvironment»\nUvicorn FastAPI\nVue dist 排程器'),'d':((90,510,600,210),'«database»\nPostgreSQL 18'),'e':((1050,940,650,150),'«external» Gemini與來源網站')},[('u','p','HTTPS'),('p','a','HTTP / WebSocket'),('a','d','SQL'),('a','e','HTTPS')],1140)
flow('packages',{'v':((70,60,730,190),'«package» frontend/src\nviews components composables'),'api':((1000,60,730,190),'«package» frontend/services\nAxios API client'),'r':((1000,430,730,190),'«package» backend/app/routers'),'s':((1000,800,730,190),'«package» backend/app/services'),'m':((70,800,730,190),'«package» models 與 core\n資料庫 設定 排程'),'c':((70,430,730,190),'«package» crawlers\n各平台解析與轉換')},[('v','api',''),('api','r','HTTP'),('r','s',''),('s','m',''),('r','c','任務啟動')],1100)
flow('components',{'u':((70,80,640,200),'«component» Vue SPA\n查詢與操作介面'),'a':((1060,80,650,200),'«component» FastAPI\nJSON XLSX Cookie'),'p':((1060,510,650,200),'«component» PostgreSQL\n持久資料與跨行程狀態'),'g':((70,510,640,200),'«component» Gemini adapter\n外部文字與向量服務')},[('u','a','REST / WebSocket'),('a','p','SQLAlchemy'),('a','g','服務呼叫')],800)
flow('lock_state',{'a':((80,70,650,170),'可取得','state'),'b':((1050,70,650,170),'持有有效租約','state'),'c':((1050,460,650,170),'到期或被接手','state'),'d':((80,460,650,170),'舊工作失效並停止','state')},[('a','b','原子取得'),('b','a','owner相符釋放'),('b','c','到期或重設'),('c','d','續租失敗'),('c','a','新取得者重新競爭')],770)
flow('article_state',{'a':((80,50,640,180),'來源資料取得','state'),'b':((1070,50,640,180),'相關且已入庫','state'),'c':((1070,450,640,180),'情緒與向量可用','state'),'d':((80,450,640,180),'正文更新\n舊分析失效','state')},[('a','b','過濾及去重'),('b','c','評分與建向量'),('c','d','重新爬取且內容變更'),('d','b','保存新內容')],780)
flow('er_content',{'p':((30,60,490,190),'platforms\nPK id\nUQ name'),'b':((650,60,490,190),'boards\nPK id\nFK platform_id'),'a':((650,470,490,250),'articles\nPK id UQ unique_id\nFK platform_id\nFK board_id author_id'),'au':((1280,60,490,190),'authors\nPK id\nUQ username'),'c':((30,910,720,200),'comments\nPK id FK article_id\nON DELETE CASCADE'),'v':((1030,910,720,200),'article_chunks\nPK article_id + chunk_index\nFK article_id CASCADE')},[('p','b','1 對 多'),('b','a','1 對 多'),('au','a','1 對 多'),('a','c','1 對 多'),('a','v','1 對 多')],1180)
flow('er_users',{'p':((600,30,600,170),'plans\nPK code'),'u':((600,350,600,190),'users\nPK id FK plan_code'),'h':((30,790,540,230),'analysis_history\nusage_counters\nFK user_id CASCADE'),'w':((630,790,540,230),'watch_keywords\nalerts\nFK user_id CASCADE'),'a':((1230,790,540,230),'audit_logs\nFK actor_id\nSET NULL')},[('p','u','1 對 多 RESTRICT'),('u','h','1 對 多'),('u','w','1 對 多'),('u','a','1 對 多')],1110)
flow('navigation',{'l':((70,50,650,170),'首頁與登入','state'),'d':((1040,50,650,170),'Dashboard','state'),'q':((1040,390,650,170),'文章詳情 問答 報表','state'),'h':((70,390,650,170),'歷史 監測 個人資料','state'),'a':((600,790,650,170),'帳號 系統 爬取管理','state')},[('l','d','登入或訪客'),('d','l','完成登出'),('d','q','功能入口'),('d','h','導覽列'),('d','a','管理員入口')],1050)
# Plan Gantt, explicitly a proposed schedule rather than an invented actual record.
im,d=canvas(920); months=['115/2','3','4','5','6','7','8','9','10','11','12','116/1']
for i,m in enumerate(months):txt(d,(480+i*103,50),m,29,100)
for r,(label,start,end) in enumerate([('需求與選題',0,2),('分析與設計',1,4),('資料庫與爬蟲',3,6),('儀表板與問答',4,7),('權限與整合',6,8),('測試與修正',7,10),('文件與交付',8,11)]):
    y=130+r*100;txt(d,(25,y+20),label,36,380,False)
    for c in range(12):d.rectangle((430+c*103,y,533+c*103,y+65),outline='#bbb',width=1)
    d.rectangle((434+start*103,y+6,529+end*103,y+59),fill='#b9c8d9',outline='#555',width=2)
txt(d,(900,850),'規劃時程；實際完成日期以版本紀錄與團隊確認為準',32,1700)
im.save(FIG/'gantt.png')

exec((WORK/'enhance_diagrams.py').read_text(encoding='utf-8'))
exec((WORK/'ooad_diagrams.py').read_text(encoding='utf-8'))
doc=Document(WORK/'reference.docx')
body=doc._element.body
for node in list(body):
    if node.tag!=qn('w:sectPr'):body.remove(node)
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(1.5)
sec.header_distance=sec.footer_distance=Cm(1)
for grid in list(sec._sectPr.findall(qn('w:docGrid'))):sec._sectPr.remove(grid)
for part in (sec.header,sec.footer):
    for child in list(part._element):part._element.remove(child)

def font_style(style,size,bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'標楷體')
    for att in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+att),None)
    pp=style.element.get_or_add_pPr()
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)

for name,size,bold in [('Normal',14,False),('Title',24,True),('Subtitle',18,False),('Heading 1',18,True),('Heading 2',16,True),('Heading 3',14,True),('Caption',14,False)]:
    if name not in doc.styles:doc.styles.add_style(name,1)
    font_style(doc.styles[name],size,bold)
    pf=doc.styles[name].paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(6);pf.line_spacing=1.0
    for el in list(doc.styles[name].element.get_or_add_pPr()):
        if el.tag in [qn('w:pBdr'),qn('w:shd')]:doc.styles[name].element.get_or_add_pPr().remove(el)
normal=doc.styles['Normal'].paragraph_format
normal.first_line_indent=Pt(28);normal.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2','Heading 3']:
    pf=doc.styles[name].paragraph_format;pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.space_before=Pt(12)
    pf.left_indent=Pt(0)
    pp=doc.styles[name].element.get_or_add_pPr()
    old=pp.find(qn('w:numPr'))
    if old is not None:pp.remove(old)
    np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);pp.append(np)
    level=OxmlElement('w:outlineLvl');level.set(qn('w:val'),str(int(name[-1])-1));doc.styles[name].element.get_or_add_pPr().append(level)
doc.styles['Heading 1'].paragraph_format.page_break_before=True
for name in ['Figure Caption','Table Caption','Code','Front Heading']:
    if name not in doc.styles:doc.styles.add_style(name,1)
    font_style(doc.styles[name],14 if name!='Front Heading' else 18,name=='Front Heading')
    pf=doc.styles[name].paragraph_format;pf.first_line_indent=Pt(0);pf.space_after=Pt(6);pf.line_spacing=1.0
doc.styles['Figure Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Table Caption'].paragraph_format.keep_with_next=True
doc.styles['Front Heading'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for i in range(1,4):
    name=f'TOC {i}'
    if name not in doc.styles:doc.styles.add_style(name,1)
    font_style(doc.styles[name],14)
    doc.styles[name].paragraph_format.first_line_indent=Pt(0)
    doc.styles[name].paragraph_format.space_after=Pt(3)

def para(t='',style=None):
    return doc.add_paragraph(t,style)
def field(p,instruction):
    run=p.add_run(); el=OxmlElement('w:fldSimple');el.set(qn('w:instr'),instruction)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='更新中';r.append(t);el.append(r);run._r.addnext(el)

# Cover, matching reference's stacked university/department/title block.
for t,size in [('國立臺北商業大學',24),('資訊管理系',22),('115年資訊系統專案設計',20)]:
    p=para(t);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(22);p.runs[0].font.size=Pt(size)
p=para('系統手冊','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
p=para('MeBOD 醫美時尚輿情分析系統','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
p=para('Medical Beauty Opinion Dashboard');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(30)
# Use actual project logo captured from real rendered SVG, if available.
logo=WORK/'screens/logo.png'
if logo.exists():
    p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.add_run().add_picture(str(logo),width=Cm(3))
else:
    p=para('MeBOD');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.runs[0].font.size=Pt(36)
for t in ['組　　別：第115405組','題　　目：MeBOD 醫美時尚輿情分析系統','指導老師：＿＿＿＿＿＿＿＿','組　　長：＿＿＿＿＿＿＿＿','組　　員：＿＿＿＿＿＿＿＿','　　　　　＿＿＿＿＿＿＿＿']:
    p=para(t);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Cm(3);p.paragraph_format.space_after=Pt(12)
p=para('中華民國115年9月28日');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(22)
sec.different_first_page_header_footer=True
foot=sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;foot.paragraph_format.first_line_indent=Pt(0)
foot.add_run('MeBOD 系統手冊　');field(foot,'PAGE')
for title,instr in [('目錄','TOC \\o "1-2" \\h \\z \\u'),('圖目錄','TOC \\t "Figure Caption,1" \\h \\z'),('表目錄','TOC \\t "Table Caption,1" \\h \\z')]:
    doc.add_page_break();para(title,'Front Heading');field(para(),' '+instr+' ')

table_names={'platforms':'社群平台','boards':'看板與來源目標','authors':'來源作者','articles':'文章主體','comments':'文章留言','article_chunks':'文章段落向量','users':'系統帳號','plans':'功能方案','usage_counters':'每月用量','watch_keywords':'個人監測關鍵字','alerts':'個人預警','analysis_history':'個人分析歷史','analysis_results':'共用分析快取','crawl_logs':'爬取工作紀錄','audit_logs':'操作稽核','settings':'系統設定','system_locks':'跨行程租約','rate_limit_hits':'速率限制視窗'}
field_notes={'id':'自動遞增識別','name':'名称或穩定識別','display_name':'顯示名稱','base_url':'來源網址','created_at':'建立時間','platform_id':'來源平台','board_id':'所屬看板','author_id':'來源作者','unique_id':'文章去重識別','title':'標題','content':'文字內容','url':'原文網址','push_count':'來源互動數','sentiment':'情緒分類 可為未評分','published_at':'來源發表時間','last_crawled_at':'最近爬取時間','floor':'留言順序','username':'登入帳號或作者代號','password_hash':'含鹽及次數的PBKDF2雜湊','role':'admin或user','avatar_emoji':'頭像符號','avatar_color':'頭像顏色','is_active':'1啟用 0停用','plan_code':'帳號方案','failed_login_count':'連續登入失敗次數','locked_until':'鎖定期限','password_changed_at':'密碼變更UTC時間','last_login_at':'最近登入時間','actor_id':'執行者帳號','actor_username':'執行者名稱快照','target_username':'目標帳號名稱','action':'操作代碼','detail':'詳細內容','keyword':'查詢或監測詞','analysis_type':'分析種類','days':'觀察天數','result_json':'分析JSON內容','expired_at':'快取到期時間','code':'方案代碼','max_watch_keywords':'監測組數上限','max_history_days':'查詢天數上限','allow_all_platforms':'跨平台能力旗標','monthly_qa_quota':'每月問答額度','allow_export':'Excel匯出開關','sort_order':'顯示順序','user_id':'資料擁有者','period':'年月 YYYY-MM','qa_count':'正常返回問答累計','updated_at':'最近更新時間','acquired_at':'取得租約時間','expires_at':'租約到期時間','owner':'每次取鎖UUID','key':'設定或限流識別','window_start':'固定視窗起點','count':'目前視窗次數','value':'設定字串','enabled':'監測啟用狀態','level':'warning或critical','negative_ratio':'負面比例快照','sentiment_score':'情緒分數快照','article_count':'文章數快照','is_read':'通知已讀狀態','status':'工作狀態','new_count':'新增數','skipped_count':'略過數','filtered_count':'過濾數','error_message':'失敗原因','started_at':'工作開始時間','finished_at':'工作結束時間','article_id':'所屬文章','chunk_index':'段落序號 從0開始','model':'向量模型名稱','dimensions':'向量維度','vector':'float32小端序位元組'}
chap='0';tn=0;fn=0;section='';captions=[]

def table(rows,title=None):
    global tn
    tn+=1
    caption=f'表{chap}-{tn} {title or re.sub(r"^\d+-\d+\s+", "", section)}'
    para(caption,'Table Caption');captions.append(caption)
    n=len(rows[0]);tb=doc.add_table(rows=0,cols=n);tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    widths={2:[5,13],3:[4.2,6,7.8],4:[3.4,4.5,4.6,5.5],5:[4,4,2,3.2,4.8]}.get(n,[18/n]*n)
    if rows[0][0]=='欄位名稱':widths=[4,3,1.8,1.8,7.4]
    if rows[0][0]=='編號與子欄位':widths=[7,4.5,6.5]
    for col,w in zip(tb.columns,widths):col.width=Cm(w)
    pr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',80),('bottom',80),('left',100),('right',100)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for ri,row in enumerate(rows):
        cells=tb.add_row().cells
        for ci,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            run=p.add_run(str(value));run.bold=ri==0
            if rows[0][0] in ['欄位名稱','編號與子欄位']:run.font.size=Pt(11)
            if ri==0:
                p.paragraph_format.keep_with_next=True
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF3');cell._tc.get_or_add_tcPr().append(sh)
        trpr=tb.rows[-1]._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit');trpr.append(no)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    p=para();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2)

def fig(path,title,maxh=15):
    global fn
    fn+=1;im=Image.open(path);w,h=im.size
    width=min(18,maxh*w/h)
    p=para();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Cm(width))
    caption=f'圖{chap}-{fn} {title}';para(caption,'Figure Caption');captions.append(caption)

def schema():
    tables=json.loads((WORK/'schema.json').read_text(encoding='utf-8'))
    order=list(table_names)
    tables.sort(key=lambda t:order.index(t['table']))
    table([['資料表','用途','欄位數']]+[[t['table'],table_names[t['table']],len(t['columns'])] for t in tables],'業務資料表一覽')
    for idx,t in enumerate(tables,1):
        name=t['table']
        p=para(f'8-2-{idx} {name} {table_names[name]}','Heading 3')
        rows=[['欄位名稱','資料型別','鍵值','NULL','說明與預設值']]
        for c in t['columns']:
            flags=['PK'] if c['pk'] else []
            if c['unique']:flags.append('UQ')
            if c['fk']:flags.append('FK')
            note=field_notes.get(c['name'],c['name'])
            if c['default']:note+=('；DB預設 ' if c['default']=='now()' else '；ORM預設 ')+c['default']
            for fk in c['fk']:note+=f"；FK → {fk['target']}；刪除 {fk['delete']}"
            dtype=c['type'].replace('TIMESTAMP WITHOUT TIME ZONE','timestamp').replace('VARCHAR','varchar').replace('INTEGER','integer').replace('TEXT','text')
            rows.append([c['name'],dtype,' '.join(flags) or '—','是' if c['nullable'] else '否',note])
        # Dictionary needs wider identifiers and explanations.
        table(rows,f'{name} 欄位定義')
        explanations={
            'articles':'文章的發表時間、入庫時間與最近爬取時間分別表示來源事件、資料建立與同步進度。相同文章重新取得時應更新內容及互動資訊，而不是建立第二個識別。',
            'users':'role與plan_code分別處理維運權限與功能額度。修改密碼後以password_changed_at比對JWT簽發時間，舊權杖會失效；密碼不以明文儲存。',
            'analysis_history':'每筆歷史保存使用者及分析JSON快照。服務查詢需加上目前使用者條件，刪除歷史不能刪除analysis_results共用快取。',
            'analysis_results':'本表為共用分析快取，不含使用者擁有者。服務以關鍵字、分析類型與時間範圍管理使用；它不能作為個人歷史的替代資料來源。',
            'article_chunks':'複合主鍵避免同一文章段落重複。vector長度需符合dimensions乘以4位元組；正文或模型變更時，應讓舊段落失效並重新產生。',
            'system_locks':'owner是取得鎖時的唯一租約識別。續租、一般釋放和過期接手需檢查條件；管理員強制釋放是另外的維運操作。',
            'usage_counters':'計次週期採年月，月份變更後建立新的計數資料。當前未在模型宣告user_id加period複合唯一約束，擴大併發使用前應補足一致性設計。',
            'audit_logs':'刪除執行者帳號後，actor_id可成為NULL，但actor_username文字仍保留，避免失去稽核可讀性。不要在detail記錄密碼或API金鑰。',
        }
        para(explanations.get(name,f'本表由對應服務維護。直接修改資料前應確認外鍵與服務規則；輸入資料仍須經API驗證，不能只依欄位型別判斷業務上是否有效。'))
    table([['資料表','欄位','用途'],['alembic_version','version_num','記錄目前遷移版本；由Alembic維護，非ORM業務資料表']],'資料庫版本管理')

text=(WORK/'manual.md').read_text(encoding='utf-8')
# Normalize stray simplified glyphs in authored prose, without touching source files.
for a,b in {'应':'應','结':'結','项':'項','号':'號','边':'邊','则':'則','为':'為','与':'與','条':'條','图':'圖','这':'這','库':'庫','说':'說','认':'認','围':'圍','确':'確','据':'據','测':'測','试':'試','签':'簽','载':'載','内':'內','页':'頁','补':'補','后续':'後續','范围':'範圍','参数':'參數','权限':'權限','回顾':'回顧','相关':'相關','图表':'圖表','定义':'定義','名称':'名稱','云':'雲'}.items():text=text.replace(a,b)
lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('# '):
        title=line[2:];m=re.search(r'第(\d+)章',title);chap=m.group(1) if m else '附';tn=fn=0
        para(title,'Heading 1');section=title;i+=1;continue
    if line.startswith('## '):section=line[3:];para(section,'Heading 2');i+=1;continue
    if line.startswith('### '):section=line[4:];para(section,'Heading 3');i+=1;continue
    if line.startswith('!FIG '):
        _,name,title=line.split(' ',2);fig(FIG/(name+'.png'),title,22 if name.startswith('seq_') else 15);i+=1;continue
    if line.startswith('!OVERVIEW '):
        _,name,title=line.split(' ',2)
        overview=doc.add_section(WD_SECTION_START.NEW_PAGE)
        overview.page_width=Cm(42);overview.page_height=Cm(29.7)
        overview.different_first_page_header_footer=False
        for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(overview,side,Cm(1.5))
        fn+=1
        p=para();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FIG/(name+'.png')),width=Cm(39))
        caption=f'圖{chap}-{fn} {title}';para(caption,'Figure Caption');captions.append(caption)
        portrait=doc.add_section(WD_SECTION_START.NEW_PAGE)
        portrait.page_width=Cm(21);portrait.page_height=Cm(29.7);portrait.different_first_page_header_footer=False
        i+=1;continue
    if line.startswith('!SCREEN '):
        _,name,title=line.split(' ',2);fig(WORK/'screens'/(name+'.png'),title,14);i+=1;continue
    if line.startswith('!SOURCE '):
        _,path,start,end=line.split();start=int(start);end=int(end)
        source=(WORK.parent.parent/path).read_text(encoding='utf-8').splitlines()
        source_p=para(f'原始碼來源：{path}，第 {start} 至 {end} 行。')
        source_p.paragraph_format.keep_with_next=True
        for lineno in range(start,end+1):
            raw=f'{lineno:>3}  '+source[lineno-1]
            pieces=[];part='';units=0
            for ch in raw:
                span=2 if unicodedata.east_asian_width(ch) in ['W','F'] else 1
                if units+span>98:pieces.append(part);part='     ';units=5
                part+=ch;units+=span
            pieces.append(part)
            p=para('\n'.join(pieces),'Code')
            p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
            p.paragraph_format.keep_with_next=lineno<end
            for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
        i+=1;continue
    if line=='!SCHEMA':schema();i+=1;continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            p=para(lines[i],'Code');p.paragraph_format.keep_with_next=i+1<len(lines) and not lines[i+1].startswith('```');i+=1
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',v) for v in vals):rows.append(vals)
            i+=1
        table(rows);continue
    para(line);i+=1

# Cache fields will be refreshed by native Word before export.
settings=doc.settings.element
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
doc.core_properties.title='MeBOD 醫美時尚輿情分析系統 系統手冊'
doc.core_properties.subject='115年資訊系統專案設計'
doc.core_properties.author='MeBOD 專題團隊'
doc.core_properties.comments=''
doc.save(OUT)
(WORK/'captions.json').write_text(json.dumps(captions,ensure_ascii=False,indent=2),encoding='utf-8')
ref=Path('C:/Users/giato/Downloads/115年_系統手冊 (3).doc')
(WORK/'reference.sha256').write_text(hashlib.sha256(ref.read_bytes()).hexdigest())
print(str(OUT));print('Figures',sum(x.startswith('圖') for x in captions),'Tables',sum(x.startswith('表') for x in captions))
