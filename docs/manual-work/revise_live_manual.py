from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

W=Path(__file__).parent
D=Document(W.parent/'MeBOD_115年系統手冊_大學部物件導向修訂版.docx')
OUT=W.parent/'MeBOD_115年系統手冊_網站操作與A4修訂版.docx'
def find(t): return next(p for p in D.paragraphs if p.text==t)
def add(anchor,text='',style='Normal'):
    p=anchor.insert_paragraph_before(text,style)
    return p
def remove_between(start,end):
    el=start._p.getnext()
    while el is not end._p:
        nxt=el.getnext();el.getparent().remove(el);el=nxt
def paras(anchor,items):
    for x in items:add(anchor,x)
def steps(anchor,items):
    for i,x in enumerate(items,1):
        p=add(anchor,f'步驟{i}：{x}')
        p.paragraph_format.keep_together=True
def section(anchor,title):add(anchor,title,'Heading 2')
def sub(anchor,title):add(anchor,title,'Heading 3')
fig_num=1
def screen(anchor,name,title):
    global fig_num
    fig_num+=1
    p=add(anchor);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(W/'live-web-screens'/f'{name}.png'),width=Cm(18))
    add(anchor,f'圖12-{fig_num} {title}','Figure Caption')

# Use A4 throughout; preserve the overview's native resolution and aspect ratio.
for s in D.sections:
    s.orientation=WD_ORIENT.PORTRAIT;s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(1.5);s.right_margin=Cm(1.5)
for shape in D.inline_shapes:
    if shape.width>Cm(18):
        ratio=Cm(18)/shape.width;shape.height=int(shape.height*ratio);shape.width=Cm(18)
for p in D.paragraphs:
    if '以下總圖在單一A3' in p.text:
        p.text=p.text.replace('單一A3橫向頁','A4直向版面')
    if '資料庫總覽保留A3' in p.text:
        p.text='格式核對：全冊A4直向；中文標楷體，英文Times New Roman；章18點、節16點、內文及目錄14點；圖名在圖下、表名在表上。圖8-1依版心等比例縮放，完整欄位與外鍵名稱可續讀分領域圖及資料字典。'
# The two section separators existed solely for the A3 foldout.
first_section=D.sections[0]._sectPr
last_section=D.sections[-1]._sectPr
for element in first_section:
    if element.tag in [qn('w:footerReference'),qn('w:headerReference')]:
        last_section.insert(0,deepcopy(element))
    elif element.tag==qn('w:titlePg'):
        last_section.append(deepcopy(element))
for p in list(D.paragraphs):
    pr=p._p.pPr
    if pr is not None and pr.find(qn('w:sectPr')) is not None:
        pr.remove(pr.find(qn('w:sectPr')))
        if not p.text and not p._p.xpath('.//w:drawing'):p._p.getparent().remove(p._p)

# Preserve audited component inventories but replace code listings with contracts.
ch9=find('第9章 程式');ch10=find('第10章 測試模型')
tables=[];el=ch9._p.getnext()
while el is not ch10._p:
    if el.tag==qn('w:tbl'):tables.append(deepcopy(el))
    el=el.getnext()
remove_between(ch9,ch10)
add(ch10,'本章依「9-1 元件清單及其規格描述」及「9-2 其他附屬之各種元件」組織，說明程式責任、介面、處理順序、相依性與例外。完整原始碼由專案目錄交付，不在手冊逐檔刊載；設計互動見第6章，元件相依見第7章，驗證案例見第10章。〔AI02〕')
section(ch10,'9-1 元件清單及其規格描述')
add(ch10,'後端檔案以下以backend/app為根目錄；前端頁面位於frontend/src/views。控制層負責驗證與HTTP回應，服務層負責業務流程，ORM與Session負責資料存取；前端顯示狀態不能取代後端授權。')
for n,title in enumerate(['後端主要元件及介面','主要HTTP介面規格','前端頁面元件'],1):
    add(ch10,f'表9-{n} {title}','Table Caption');ch10._p.addprevious(tables[n-1]);add(ch10)
add(ch10,'介面補充：GET /api/dashboard/full與GET /api/articles/{article_id}為公開讀取；文章不存在回404。私人歷史需登入並依user_id過濾。Excel另檢查方案及可查天數。完整請求型別以目前版本OpenAPI與路由驗證為準。')
specs=[
('9-1-1 身分驗證與工作階段元件',[
'目的與介面：auth路由及auth_service處理登入、目前帳號、登出與受保護操作的身分驗證。輸入為帳號、密碼或Cookie；輸出為帳號顯示資訊、登入Cookie或明確HTTP錯誤。密碼與JWT不得作為一般日誌內容。',
'處理順序：先套用登入頻率限制，再檢查帳號狀態與密碼；成功時產生JWT並設定HttpOnly Cookie。受保護請求重新驗證JWT、帳號狀態與密碼變更時間，再判斷管理員角色。登出清除Cookie；前端同步清除顯示狀態。',
'例外與驗證：錯誤憑證回401，禁止的角色或方案回403，鎖定與頻率限制依端點回423或429。驗證成功登入、me、登出、停用帳號與密碼變更後舊JWT失效。依賴users、auth設定與限流儲存。']),
('9-1-2 儀表板、文章與個人歷史元件',[
'目的與輸入：dashboard_service接受關鍵字、天數、平台／看板與排序條件，輸出統計、趨勢、情緒、熱門文章及洞察所需JSON。前端useDashboard管理載入、錯誤與篩選狀態；ArticleDetailView依文章id顯示正文及已收錄留言。',
'查詢程序：正規化平台與看板條件，建立一致查詢範圍，取得文章與聚合結果，再轉換為前端資料結構。儀表板的行程內快取期限為10分鐘；快取不是跨行程共享資料庫，也不代表來源網站即時更新。',
'歷史程序：關鍵字分析允許可選身分；已登入者才建立個人歷史。讀取與刪除以目前user_id限制，跨帳號或不存在資料刪除回404。訪客History應呈現登入提示，不送出私人歷史請求。',
'規格限制：歷史紀錄不保證所有指標都完全凍結，部分值可能依目前資料重新計算；dashboard/full目前沒有方案天數裁切。驗證同名跨平台看板、空資料、排序、快取條件與帳號隔離。']),
('9-1-3 AI問答與檢索元件',[
'介面與前置條件：POST /api/qa/ask需登入並通過方案額度與每分鐘10次限制。question長度2至500字；歷史訊息角色限制user或assistant，每筆內容最多2000字。輸出包含回答及可用來源資訊。',
'主要程序：先檢查額度，再以問題、完整Dashboard脈絡及最近六筆對話建立SHA256快取識別；15分鐘內命中可直接回覆。未命中時，有Dashboard脈絡就直接依脈絡生成；沒有脈絡才走關鍵字與向量混合檢索，再交模型回答。',
'檢索規格：embedding_service將文章切段，段長400字、重疊80字、每篇最多30段；768維float32向量以bytea保存，並非pgvector欄位。混合排名使用RRF常數60。正文或內容更新時應清除過期段落後補建。',
'計次與例外：路由在回答函式正常返回後記錄用量；快取回答仍計次，正常返回的降級回答也可能計次。不得把「未再次呼叫模型」解讀成「不計額度」。驗證兩條生成分支、快取、空來源、模型降級及額度不足。']),
('9-1-4 爬取、文章更新與租約元件',[
'責任與相依：crawler路由接收平台、看板及頁數，驗證管理員後排入背景工作。lock_service以資料庫system_locks與UUID擁有者識別控制同時執行；article_service正規化文章、作者、看板與留言。',
'處理順序：取得租約後爬取來源，過濾不相關內容，以unique_id判斷新增或更新，保存有效欄位，再依工作流程評分、補建向量與檢查預警。執行中續租，完成或失敗時依擁有者釋放；租約被其他有效工作持有時手動請求回409。',
'資料更新規格：新內容為空時保留既有有效內容。正文（【留言】之前）改動使情緒過期；整體內容改動使向量段落過期。部分留言批次以重複計數處理累積資料，不因一次少量回傳就刪掉已收錄留言。',
'交易與恢復：文章與留言提交並非一個涵蓋全部後處理的大交易。手動流程的租約涵蓋後處理；每日流程在爬取後釋放，再進行情緒、向量與預警。失敗時先讀工作紀錄，確認已提交部分，避免把重新執行視為整批回滾。']),
('9-1-5 匯出、方案、監測與管理元件',[
'匯出規格：GET /api/export/articles.xlsx先驗證帳號及匯出權限，正規化看板篩選，依方案裁切查詢天數，再封裝XLSX。輸入days為1至365；檔名使用請求天數，不保證內容涵蓋相同天數，因此驗收須檢查實際資料日期。',
'監測規格：新增監測詞先去除前後空白、驗證長度與天數，再檢查同帳號重複及方案組數，成功後保存user_id與稽核紀錄。重複關鍵字回409；本人清單依user_id過濾。負面比例與最少文章數共同決定是否建立預警。',
'管理規格：帳號與系統設定由管理員介面提供，後端仍需require_admin驗證。帳號角色、啟用狀態、方案與排程設定各由所屬端點處理；變更後重新讀取顯示資料，並檢查適用的audit_logs。',
'現況限制與驗證：部分預警已讀異動及手動全域檢查仍有跨帳號隔離待補強事項，詳第6章與第10章，不能宣稱所有監測操作均已完整隔離。驗證自己與他人資料、方案不足、非法設定及重複請求。'])]
for title,body in specs:sub(ch10,title);paras(ch10,body)
section(ch10,'9-2 其他附屬之各種元件')
for title,body in [
('9-2-1 前端共用元件與API連線',['services/api.js集中Axios連線設定並以withCredentials攜帶Cookie；useAuth封裝登入顯示狀態，useDashboard管理查詢狀態，useAlerts管理通知，useWebSocket管理即時訊息。頁面元件應使用共用介面，避免各自形成不同授權或錯誤規則。','已登入狀態收到401時清理本機顯示資料並引導登入。訪客頁面應先判斷功能是否需要帳號。WebSocket目前為公開且行程內廣播，不能將其當成有帳號隔離或跨行程可靠佇列的通道。']),
('9-2-2 資料庫、遷移與啟停元件',['core/database.py建立Engine與Session，請求透過get_db取得會話；成功提交、例外回滾與關閉需依服務交易邊界處理。database_models定義18張ORM表，Alembic版本表另管理遷移狀態。資料表修改須同步第8章字典、遷移與測試。','core/startup.py準備預設方案、管理員及系統設定；core/scheduler.py依台灣時間排程每日任務並處理補跑；關閉時停止相關背景資源。設定需由環境與受控設定介面提供，避免將資料庫密碼或API金鑰硬編碼進原始碼。']),
('9-2-3 補建腳本、測試與建置元件',['backend/scripts/backfill_sentiment.py用於補評情緒，backfill_embeddings.py用於補建向量，compare_retrieval.py比較檢索結果。執行前確認目標資料庫、處理筆數、模型費用與重跑條件；執行後核對未完成數量及錯誤紀錄。','run_system_tests.py整合獨立PostgreSQL與瀏覽器全流程測試入口；後端單元測試、前端測試、資料庫整合測試及E2E負責不同層級。前端建置產生部署資產，不能以建置成功取代API與帳號隔離測試；實際指令及既有結果見第10、11章。']),
('9-2-4 共同錯誤合約與修改程序',['HTTP 401表示未登入或憑證失效；403表示角色、方案或設定不允許；404表示資源不存在或在目前帳號範圍不可辨識；409表示重複資料或工作衝突；422表示請求格式驗證失敗；429表示頻率超限；503表示依賴暫不可用。前端呈現可理解訊息，維運日誌保留追查資訊。','修改流程：先確認元件輸入輸出與權限，再調整服務或路由；若改資料結構，同步ORM及Alembic；若改回應欄位，同步前端使用點；最後執行成功、錯誤與跨帳號案例。第6章合約、第7章元件圖、第8章字典及本章規格須隨程式共同更新。'])]:sub(ch10,title);paras(ch10,body)

# Retain the actual OO navigation diagram, replacing only the operational screenshots/text.
ch12=find('第12章 使用手冊');ch13=find('第13章 感想')
nav_image=deepcopy(find('圖12-1 主要畫面轉移圖')._p.getprevious())
remove_between(ch12,ch13)
add(ch13,'本章以https://mebod.clouda.dpdns.org/的實際部署介面說明操作。網站截圖擷取於2026年9月29日，採桌面1440 × 900視窗設定；數量、日期與帳號可見功能會隨資料及權限改變。圖中為部署資料，不能視為固定測試預期值。〔AI02〕')
section(ch13,'12-1 畫面轉移與使用順序')
ch13._p.addprevious(nav_image);add(ch13,'圖12-1 主要畫面轉移圖','Figure Caption')
paras(ch13,['一般操作依序為「首頁 → 登入或訪客 → Dashboard → 設定條件 → 閱讀統計與文章 → AI問答或報表」。已登入帳號可從頂部進入History、監控預警及個人帳號；管理員另可進入帳號管理與系統管理。','Dashboard左側的總覽、趨勢、情緒、熱門文章、洞察與文字雲用於定位同一頁的區塊；爬蟲管理另有工作頁。文章詳情與報表以「← 返回」回到先前畫面。切換頁面後先等資料載入完成，再判讀空清單或零值。'])
section(ch13,'12-2 首頁、登入與訪客模式')
screen(ch13,'landing','正式網站產品首頁（已登入狀態）')
steps(ch13,['開啟網站，閱讀「功能特色」「運作流程」「適用對象」與「常見問題」。已登入時首頁顯示「進入系統」「前往Dashboard」；未登入時顯示登入或體驗入口。','需要個人功能時進入登入頁，輸入管理員提供的帳號與密碼，按「登入」。成功後確認頂端帳號名稱與Dashboard已出現；帳密錯誤時停留原頁更正，不要連續快速重試。','只體驗公開查詢時選「以訪客身分瀏覽Dashboard」。訪客可閱讀公開統計與文章，但私人歷史、問答、監控及匯出仍依帳號／方案限制。需要這些功能時改以正式帳號登入。','工作結束按頂部「登出」，確認返回登入流程。遇到登入失效提示時重新登入後再執行原操作。'])
section(ch13,'12-3 Dashboard搜尋與平台篩選')
screen(ch13,'dashboard','Dashboard總覽與搜尋條件')
steps(ch13,['先確認現有關鍵字標籤。點標籤上的「×」移除不需要的詞，或在「輸入其他關鍵字後按Enter新增」輸入新詞並按Enter。確認標籤已加入，避免只打字卻未加入條件。','從下拉選單選近7天、近30天或近90天。點「全部」或PTT、Dcard、Mobile01、Threads設定平台，再按「搜尋」。查詢後核對關鍵字、天數與文章來源。','查看資料狀態及更新時間，再讀相關文章數、平均推文數、正面比例與負面預警數。沒有文章時先放寬天數或移除過多條件；仍無資料時到爬蟲管理確認來源是否已收錄。'])
screen(ch13,'search-advanced','展開進階搜尋與建議關鍵字')
steps(ch13,['按「進階搜尋」展開更多關鍵字與最近搜尋，點欲使用的詞加入條件；使用「收合」回到精簡版面。','需要重新設定時使用「清除條件」，再檢查目前標籤與平台後搜尋。不要將不同平台的互動數直接當成同一尺度的情緒分數。'])
section(ch13,'12-4 趨勢、情緒與洞察閱讀')
screen(ch13,'sentiment','情緒分佈、熱門話題與文章區塊')
steps(ch13,['按左側「趨勢」定位時間序列，對照同一查詢期間的討論量變化；尖峰只是收錄資料中的變化，應再回查當日文章。','按「情緒」查看正面、中性、負面占比，再查看熱門話題。注意畫面部分指標標示為依推文數粗略估計；AI覆蓋率表示已評分比例，不是準確率。','按「洞察」依序閱讀摘要、熱門話題、消費者痛點與行銷建議。建議以根據、做什麼、為什麼、誰來做、在哪裡、何時、怎麼做、投入及預期成效呈現，應再用原文驗證其依據。','按「文字雲」查看高頻詞。詞頻代表出現次數，不等同不重複人數、購買意圖或因果關係；可將值得追蹤的詞帶回搜尋區進一步分析。'])
screen(ch13,'insights','LLM洞察與行銷建議閱讀區')
section(ch13,'12-5 熱門文章、詳情與來源核對')
screen(ch13,'articles','熱門文章清單與排序')
steps(ch13,['按「熱門文章」定位列表，依目的選「推文數」「最新」或「相關度」排序。先看來源、日期及摘要，再選擇要追查的文章。','點該列「詳情」進入站內文章頁，核對標題、作者、時間、互動數及正文。需要確認外部內容時使用原文連結。畫面的通用PTT提示不代表每篇都來自PTT，應以該列來源及實際連結為準。'])
screen(ch13,'article-detail','Threads文章詳情與原文入口')
paras(ch13,['若有收錄留言，續讀留言內容及其情緒統計；本圖所選文章沒有留言清單，不能將沒有收錄解讀為來源網站完全沒有留言。原文被刪除或限制存取時，站內保存內容可能與目前原站不同。完成後按「← 返回」。'])
section(ch13,'12-6 AI問答與對話管理')
screen(ch13,'qa','AI問答與Dashboard脈絡提示')
steps(ch13,['先在Dashboard完成主題與期間查詢，再按頂部「AI問答」。閱讀提示中目前依據的Dashboard關鍵字，確認不是先前任務的脈絡。','點建議問題，或在輸入框輸入具體問題，例如「根據目前Dashboard，最大的負面風險是什麼？」。按Enter或「送出問題」；Shift+Enter可換行。','等待回答後，將重點、行動建議及可用來源與文章核對。沒有足夠來源時調整問題或回Dashboard修改查詢，不把模型回答視為已查證的醫療或營運結論。','要切換工作主題時使用左側「新對話」，已有對話可透過搜尋框尋找。額度或功能不足時依畫面訊息處理；重新送出問題可能再次計次。'])
section(ch13,'12-7 個人歷史、比較與趨勢')
screen(ch13,'history','個人分析歷史清單')
steps(ch13,['按頂部「History」，確認顯示自己的分析紀錄。用「搜尋關鍵字」收斂資料；需要追蹤的重要紀錄可加星並用「只看已加星」篩選。加星保存在目前瀏覽器，不保證跨裝置同步。','勾選兩筆紀錄後按「前往比較」或「並排比較」，核對兩側實際日期、關鍵字及期間，再閱讀文章數、情緒分數與負面比例的差異。','本次部署畫面中「上次分析」「本次分析」標籤與日期順序可能不一致，應以實際時間判斷前後。相同指標不代表資料完全相同，也不能推定全部欄位為當時凍結快照。'])
screen(ch13,'compare','兩筆歷史的並排比較')
paras(ch13,['選「趨勢圖」可繼續觀察歷次分析變化。需要刪除時先核對該筆日期與關鍵字，再使用「刪除」並確認；刪除個人歷史不等於刪除來源文章。訪客應看到登入提示，登入失效則重新登入。'])
section(ch13,'12-8 Excel下載與列印報表')
screen(ch13,'report','正式網站的可列印分析報告')
steps(ch13,['回Dashboard核對搜尋條件後按「匯出Excel」。此操作需登入及匯出權限；下載完成後以試算表程式開啟.xlsx，檢查文章內容、來源與實際日期。','匯出受方案天數及筆數限制，檔名中的天數可能保留原請求值；若筆數或期間少於畫面預期，先核對方案與實際日期，不只看檔名。','按「匯出報告」，等待「報告產生中」消失，確認關鍵字、期間、產生時間、輿情總覽、摘要與熱門文章。','按「列印／存成PDF」開啟瀏覽器列印預覽，選目的地與紙張後檢查分頁，再儲存或列印。報表頁提供閱讀整理，Excel提供後續資料處理。'])
section(ch13,'12-9 監控關鍵字與風險預警')
screen(ch13,'monitor','監控關鍵字與未讀風險預警')
steps(ch13,['按「監控預警」，在「輸入要監控的關鍵字」填入品牌或議題，選觀察天數後按「加入」。成功後確認新詞出現在清單；重複或超過方案限制時依訊息更正。','使用該詞的狀態按鈕切換監控狀態；移除前確認是否仍需長期追蹤。管理員畫面可能顯示不受額度限制，一般帳號仍應依實際方案操作。','在「風險預警」閱讀警示關鍵字、統計期間、文章數、負面比例與時間，再返回Dashboard查詢同詞及原文。按「立即檢查」會執行檢查；「執行每日任務」屬管理維運操作，可能觸發實際處理工作。','完成閱讀後可按單筆「標為已讀」；「全部已讀」會改變通知狀態。已讀只表示看過通知，不表示負面議題已解決。'])
section(ch13,'12-10 個人帳號與密碼')
screen(ch13,'profile','帳號資訊與密碼表單')
steps(ch13,['點頂部自己的顯示名稱，核對帳號、角色、啟用狀態、最後登入及建立時間。點名稱或頭像旁的編輯按鈕進入對應編輯。','變更密碼時填舊密碼、新密碼與確認新密碼，依畫面至少8字元規則完成後按「更新密碼」。輸入不一致或舊密碼錯誤時更正後再送出。','成功後依提示重新登入；密碼變更會使舊JWT失效，其他裝置也可能需重新驗證。離開共用電腦前登出。'])
section(ch13,'12-11 管理員帳號管理')
screen(ch13,'admin-users','管理員帳號建立、篩選與管理')
steps(ch13,['以管理員登入後開啟「帳號管理」。先查看總帳號數、管理員數及啟用狀態，再以搜尋欄、角色及狀態選單找到目標帳號。','新增時依表單填帳號、顯示名稱、密碼及角色，核對權限後按「建立帳號」。建立表單目前標示至少6字元，個人改密碼頁標示至少8字元，應依各表單當下驗證規則處理並使用較強密碼。','既有帳號使用「編輯」管理允許欄位，儲存後重新核對清單；如需停用或刪除，先確認資料保留與權限影響。畫面對目前登入者的刪除按鈕有停用保護。','向下查看「操作紀錄（稽核log）」確認相關異動。交付一般帳號前，用其實際權限驗證功能，不以管理員可見功能推定所有使用者皆可使用。'])
section(ch13,'12-12 系統設定、排程與看板')
screen(ch13,'system-settings','系統設定中的預警門檻與排程')
steps(ch13,['進入「系統管理 → 系統總覽」，先看資料庫連線、API設定狀態、文章數、AI覆蓋率、排程與最後爬取。安全性提醒若存在，依其內容處理；連線正常不代表排程已在預期時間成功執行。','切到「系統設定」，記錄警示門檻、危機門檻及最少文章數。調整前確認門檻單位為負面比例百分比，且樣本數太少時不預警。','設定每日爬取啟用狀態、台灣時間的小時、每看板頁數、啟動補跑及各瀏覽器平台開關，核對後按「儲存設定」。重新讀取確認已保存，並在下一輪工作檢查結果。','「看板管理」用於查看平台收錄設定，「操作紀錄」用於追查管理異動。瀏覽器型來源需可用Chromium及來源站允許存取，僅勾選開關不代表必然成功。'])
section(ch13,'12-13 爬蟲工作、進度與執行紀錄')
screen(ch13,'crawl','爬蟲管理的工作狀態與手動表單')
steps(ch13,['從Dashboard的「前往爬蟲管理」開啟工作頁，查看執行狀態、上次爬取時間與新增／跳過／過濾數。日期過舊時先查排程及日誌。','手動工作先選平台，再選對應看板、頁數與時間範圍；確認伺服器、來源限制與處理成本後按「開始爬取」。已有工作時不要反覆點擊，後端會檢查租約。','在「即時進度」查看開始時間、目前頁數、新增數及狀態。進度停止時先檢查連線，再到「爬蟲執行紀錄」重新整理，依成功／失敗篩選並點「檢視」。','閱讀錯誤訊息後再決定重試範圍。新增表示新識別文章，跳過重複不保證所有欄位完全未更新，過濾表示未收錄。未評分資料可依權限使用「重新評分情緒」，並檢查完成後的覆蓋數。'])
section(ch13,'12-14 常見操作問題與完成檢查')
find('12-14 常見操作問題與完成檢查').paragraph_format.page_break_before=True
paras(ch13,['看到401：訪客要使用私人功能時先登入；已登入者重新驗證。看到403：檢查帳號角色、方案與平台設定。看到429：等待限制窗口後重試，避免反覆送出。','查詢沒有資料：核對關鍵字標籤是否已加入、期間是否太窄、平台是否正確及最後收錄時間。分析缺少來源或摘要時，回到文章確認資料完整性。','比較、圖表或報告數字不一致：先統一關鍵字、天數、平台與產生時間，再檢查歷史動態計算及匯出方案限制。AI覆蓋率、互動數、情緒比例與模型準確率是不同概念。','完成一次分析後，確認搜尋條件、主要指標、代表文章及報表條件均相符；重要結論保留原文參照，個人工作完成後登出。'])

# Update project-specific explanatory references without altering unrelated chapters.
for p in D.paragraphs:
    if p.style.name.startswith('toc'):continue
    for old,new in [('本章6-2與6-3','本章6-1-2與6-1-3'),('9-3至9-7','9-1與9-2'),('12-10及12-11','12-11及12-12'),('12-10 及 12-11','12-11 及 12-12')]:
        if old in p.text:p.text=p.text.replace(old,new)
ref=find('14-1 人工智慧輔助使用說明')
add(ref,'MeBOD正式部署網站。https://mebod.clouda.dpdns.org/。介面觀察與截圖日期：2026年9月29日；第12章使用者操作說明之畫面來源。')
D.save(OUT)
print(OUT)
print('Screenshots',fig_num-1,'Figures',len(D.inline_shapes),'Tables',len(D.tables),'Sections',len(D.sections))
