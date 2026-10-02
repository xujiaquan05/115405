from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import math,json

W=Path(__file__).parent; F=W/'polished-uml'; F.mkdir(exist_ok=True)
D=Document(W/'diagram-polish-base.docx')
OUT=W.parent/'MeBOD_115年系統手冊_UML圖面優化版.docx'
INK='#23364D'; BLUE='#EAF2FA'; TEAL='#EAF5F1'; GRAY='#F2F4F7'
def font(n):return ImageFont.truetype('C:/Windows/Fonts/msjh.ttc',n)
def canvas(h):return Image.new('RGB',(1600,h),'white')
def text(im,x,y,s,size=38,w=None,center=False):
 d=ImageDraw.Draw(im);yy=y
 for t in s.split('\n'):
  if w:assert font(size).getlength(t)<=w,(t,size,w)
  xx=x+(w-font(size).getlength(t))/2 if center else x
  d.text((xx,yy),t,font=font(size),fill=INK);yy+=size+13
 return yy
def line(im,pts,dash=False,head=None):
 d=ImageDraw.Draw(im)
 for a,b in zip(pts,pts[1:]):
  L=math.dist(a,b)
  if dash:
   for n in range(0,int(L),20):
    p=n/L;q=min(n+11,L)/L;d.line((a[0]+(b[0]-a[0])*p,a[1]+(b[1]-a[1])*p,a[0]+(b[0]-a[0])*q,a[1]+(b[1]-a[1])*q),fill=INK,width=3)
  else:d.line((a,b),fill=INK,width=3)
 if head:
  a,b=pts[-2:];L=math.dist(a,b);ux=(b[0]-a[0])/L;uy=(b[1]-a[1])/L
  p=(b[0]-24*ux-12*uy,b[1]-24*uy+12*ux);q=(b[0]-24*ux+12*uy,b[1]-24*uy-12*ux)
  if head=='triangle':d.polygon([b,p,q],fill='white',outline=INK,width=3)
  elif head=='filled':d.polygon([b,p,q],fill=INK)
  else:d.line([p,b,q],fill=INK,width=3)
def tag(im,x,y,s,size=33):
 d=ImageDraw.Draw(im);ww=max(font(size).getlength(t) for t in s.split('\n'));hh=len(s.split('\n'))*(size+13)
 d.rectangle((x-9,y-4,x+ww+9,y+hh+2),fill='white');text(im,x,y,s,size)
def box(im,r,title,body='',stereo='',fill=BLUE,size=38):
 x,y,w,h=r;d=ImageDraw.Draw(im);d.rectangle((x,y,x+w,y+h),fill='white',outline=INK,width=3)
 hh=125 if stereo else 78;d.rectangle((x+2,y+2,x+w-2,y+hh),fill=fill)
 if stereo:text(im,x+12,y+12,stereo,32,w-24,True)
 text(im,x+12,y+(60 if stereo else 16),title,40,w-24,True)
 d.line((x,y+hh,x+w,y+hh),fill=INK,width=2)
 if body:text(im,x+22,y+hh+20,body,size,w-44)
def save(im,name):im.save(F/(name+'.png'))
def note(im,y,s):text(im,35,y,s,32,1530)

# Analysis: visual groupings and conceptual collaborators, not invented code classes.
im=canvas(1050)
for y,title,body,st,col in [(35,'邊界物件','登入與查詢畫面　問答畫面　歷史與管理畫面','«boundary»',BLUE),(375,'控制物件','身分驗證　篩選分析　問答協調　歷史與監測管理','«control»',TEAL),(715,'實體物件','使用者　文章　分析歷史　監測詞　預警','«entity»',GRAY)]:
 box(im,(180,y,1240,240),title,body,st,col,38)
line(im,[(800,275),(800,375)]);tag(im,845,292,'參與協作')
line(im,[(800,615),(800,715)]);tag(im,845,632,'處理領域資料')
note(im,992,'分析責任分組；實線為協作關聯，不表示呼叫先後或資料表外鍵。');save(im,'analysis_layers')
im=canvas(1530)
for col,title in enumerate(['邊界 Boundary','控制 Control','實體 Entity']):text(im,30+col*540,10,title,38,460,True)
rows=[('UC01 登入','登入介面','登入控制','使用者'),('UC02／UC03 查詢與問答','查詢與問答介面','查詢與回答控制','文章'),('UC05 個人歷史','歷史介面','歷史管理控制','分析歷史'),('UC05 監測','監測設定介面','監測管理控制','監測詞／預警')]
for i,(lab,a,b,c) in enumerate(rows):
 y=100+i*335;text(im,35,y,lab,34)
 for j,(v,st,col) in enumerate(zip([a,b,c],['«boundary»','«control»','«entity»'],[BLUE,TEAL,GRAY])):
  box(im,(30+j*540,y+65,460,170),v,'',st,col)
  if j<2:line(im,[(490+j*540,y+150),(570+j*540,y+150)])
note(im,1460,'各列為獨立使用個案的概念協作；控制可由多個路由與服務函式實作。');save(im,'analysis_collaboration')

# Read downward, trace upward. Explicit UML trace dependencies.
im=canvas(1580)
steps=[('使用個案與需求','第5章　角色　情境　成功與例外'),('系統事件與合約','6-1-2與6-1-3　系統循序　前後置條件'),('詳細循序與設計類別','6-1-4至6-2　責任分配　協作與關聯'),('原始碼與資料模型','第8與9章　資料結構　元件實作'),('驗證案例','第10章　正常流程　權限　失敗與競爭')]
for i,(a,b) in enumerate(steps):
 y=30+i*290;box(im,(300,y,1250,215),a,b,'«artifact»',BLUE if i%2==0 else TEAL,36)
 text(im,55,y+68,f'0{i+1}',62,180,True)
 if i:line(im,[(925,y),(925,y-75)],True,'open');tag(im,1020,y-62,'«trace»',30)
note(im,1505,'由上而下閱讀設計成果；虛線箭頭向上指向被追溯的產物。');save(im,'trace')

# ORM: spacious connectors and endpoint multiplicities.
im=canvas(1510)
box(im,(40,35,570,240),'Platform','+ id: Integer\n+ name: String',fill=BLUE)
box(im,(990,35,570,240),'Board','+ id: Integer\n+ platform_id: Integer',fill=BLUE)
box(im,(500,550,600,370),'Article','+ id: Integer\n+ unique_id: String\n+ platform_id: Integer\n+ board_id: Integer',fill=TEAL)
box(im,(40,1190,650,230),'Comment','+ id: Integer\n+ article_id: Integer',fill=GRAY)
box(im,(910,1190,650,230),'ArticleChunk','+ article_id: Integer\n+ chunk_index: Integer',fill=GRAY)
line(im,[(610,155),(990,155)]);tag(im,630,98,'1');tag(im,865,98,'0..*');tag(im,715,184,'包含',30)
line(im,[(325,275),(325,660),(500,660)]);tag(im,350,290,'1');tag(im,395,600,'0..*')
line(im,[(1275,275),(1275,660),(1100,660)]);tag(im,1300,290,'0..1');tag(im,1125,600,'0..*')
line(im,[(650,920),(650,1050),(365,1050),(365,1190)]);tag(im,680,935,'1');tag(im,385,1130,'0..*')
line(im,[(950,920),(950,1050),(1235,1050),(1235,1190)]);tag(im,980,935,'1');tag(im,1255,1130,'0..*')
note(im,1460,'屬性節錄；實線為關聯，端點數字為多重性。其餘類別及欄位見第8章。');save(im,'orm')

im=canvas(1240)
box(im,(430,30,740,280),'BrowserCrawler','- _sleep()\n- _open_page(disable_cache)',fill=BLUE,size=36)
for x,title in [(20,'DcardCrawler'),(560,'Mobile01Crawler'),(1100,'ThreadsCrawler')]:
 box(im,(x,550,480,240),title,'+ crawl_board(...)',fill=TEAL,size=34)
 line(im,[(x+240,550),(x+240,420),(800,420)])
line(im,[(800,420),(800,310)],head='triangle');tag(im,850,355,'一般化',32)
box(im,(430,920,740,230),'PTTCrawler','+ crawl_board(...)',fill=GRAY)
note(im,1180,'空心三角指向父類別；PTTCrawler獨立實作，因此不與上方繼承樹相連。');save(im,'crawler')

# Collaboration sequence: larger letters, shorter messages, consistent spacing.
im=canvas(1190);xs=[140,580,1020,1460]
for x,title,st in zip(xs,['Vue邊界','FastAPI路由','業務服務','ORM／Session'],['«boundary»','«module»','«module»','«entity»']):
 d=ImageDraw.Draw(im);d.rectangle((x-130,20,x+130,170),fill=BLUE,outline=INK,width=3);text(im,x-125,38,st,29,250,True);text(im,x-125,91,title,33,250,True);line(im,[(x,170),(x,1110)],True)
for i,(a,b,s,ret) in enumerate([(0,1,'HTTP請求',False),(1,2,'委派業務操作',False),(2,3,'查詢或更新',False),(3,2,'實體與交易結果',True),(2,1,'結果或例外',True),(1,0,'JSON或檔案',True)]):
 y=275+i*140;text(im,min(xs[a],xs[b])+20,y-60,s,35,400);line(im,[(xs[a],y),(xs[b],y)],ret,'open' if ret else 'filled')
note(im,1130,'實線箭頭：請求／呼叫　　虛線箭頭：回傳　　提交交易由實作端明確決定。');save(im,'responsibility')

def node(im,r,title,body):
 x,y,w,h=r;d=ImageDraw.Draw(im);d.polygon([(x,y),(x+18,y-18),(x+w+18,y-18),(x+w+18,y+h-18),(x+w,y+h),(x+w,y)],fill='#DCE5EE',outline=INK,width=2)
 box(im,r,title,body,'«node»',BLUE,36)
im=canvas(1390)
node(im,(100,30,1400,300),'使用者裝置','«executionEnvironment» 瀏覽器\n«artifact» Vue建置產物')
node(im,(100,490,1400,310),'應用主機','«executionEnvironment» Python與Uvicorn\n«artifact» FastAPI app　背景排程與爬取')
node(im,(50,1000,710,280),'資料庫節點','PostgreSQL\n業務資料與共享租約')
node(im,(870,1000,670,280),'外部服務','Gemini API\n社群來源網站')
line(im,[(800,330),(800,490)]);tag(im,845,377,'HTTPS／WSS')
line(im,[(400,800),(400,1000)]);tag(im,440,868,'SQL連線')
line(im,[(1200,800),(1200,1000)]);tag(im,1235,868,'HTTPS')
note(im,1330,'邏輯部署節點可位於同一實體主機；TLS與反向代理依實際環境配置。');save(im,'deployment')

def package(im,r,title,body):
 x,y,w,h=r;d=ImageDraw.Draw(im);d.rectangle((x,y,x+180,y+45),fill=BLUE,outline=INK,width=3);box(im,(x,y+45,w,h-45),title,body,fill=BLUE,size=33)
im=canvas(1480)
package(im,(40,20,1520,230),'frontend/src','views　components　composables　services/api.js')
package(im,(400,430,800,200),'app.routers','HTTP協定　驗證與依賴注入')
package(im,(400,800,800,200),'app.services','分析　問答　文章保存　方案')
package(im,(40,1200,650,200),'app.models','ORM類別與關聯')
package(im,(910,1200,650,200),'app.core 共用基礎','database　config　time_utils')
line(im,[(800,250),(800,430)],True,'open');tag(im,850,300,'«use» HTTP API')
line(im,[(800,630),(800,800)],True,'open');tag(im,850,683,'«import»')
line(im,[(600,1000),(600,1090),(365,1090),(365,1200)],True,'open');tag(im,330,1020,'«import»')
line(im,[(1000,1000),(1000,1090),(1235,1090),(1235,1200)],True,'open');tag(im,1080,1020,'«import»')
note(im,1425,'主請求相依節錄；core.scheduler另協調crawlers與services，並非整包單向分層。');save(im,'packages')

im=canvas(1320)
box(im,(420,25,760,245),'Web UI','Vue頁面與狀態','«component»',BLUE)
box(im,(420,455,760,245),'API應用','FastAPI路由與服務','«component»',TEAL)
box(im,(30,920,710,260),'背景處理','排程　爬蟲　模型呼叫','«component»',BLUE,35)
box(im,(890,920,680,260),'資料存取','SQLAlchemy與PostgreSQL','«component»',GRAY,35)
line(im,[(800,270),(800,455)],True,'open');tag(im,845,330,'HTTP／WebSocket')
line(im,[(610,700),(610,790),(385,790),(385,920)],True,'open');tag(im,65,745,'啟動背景工作')
line(im,[(1000,700),(1000,790),(1230,790),(1230,920)],True,'open');tag(im,1235,790,'SQL')
line(im,[(740,1070),(890,1070)],True,'open');tag(im,757,1100,'保存',31)
note(im,1240,'虛線箭頭由使用者指向被使用元件；這是介面依賴，不是執行時間順序。');save(im,'components')

def state(im,r,title,body):
 x,y,w,h=r;d=ImageDraw.Draw(im);d.rounded_rectangle((x,y,x+w,y+h),radius=28,fill=BLUE,outline=INK,width=3);text(im,x+15,y+20,title,43,w-30,True);d.line((x,y+85,x+w,y+85),fill=INK,width=2);text(im,x+25,y+106,body,35,w-50)
im=canvas(1570);d=ImageDraw.Draw(im);d.ellipse((777,10,823,56),fill=INK);line(im,[(800,56),(800,115)],head='open')
state(im,(430,115,740,230),'無租約資料列','尚未取得，或已釋放')
state(im,(430,650,740,245),'有效租約','owner已指定；expires_at > now')
state(im,(430,1210,740,240),'已過期租約','資料列仍在；expires_at <= now')
line(im,[(700,345),(700,650)],head='open');tag(im,740,410,'try_acquire [INSERT成功]\n/ 設owner與到期時間',33)
line(im,[(430,740),(220,740),(220,230),(430,230)],head='open');tag(im,25,410,'release\n[owner相符]\n/ 刪除資料列',31)
line(im,[(800,895),(800,1210)],head='open');tag(im,835,990,'到期 [未續租]\n/ 轉為可接手',33)
line(im,[(430,1320),(160,1320),(160,840),(430,840)],head='open');tag(im,22,955,'try_acquire\n[條件更新成功]\n/ 更換owner',31)
line(im,[(1170,715),(1480,715),(1480,855),(1170,855)],head='open');tag(im,1210,550,'renew [同owner\n且未到期]\n/ 延長期限',31)
note(im,1500,'轉移：事件 [條件] / 效果。owner不符不刪鎖；強制釋放為管理員維運例外。');save(im,'lock_state')

im=canvas(1510);d=ImageDraw.Draw(im);d.ellipse((777,10,823,56),fill=INK);line(im,[(800,56),(800,110)],head='open')
state(im,(430,110,740,200),'尚未入庫','文章主體尚未保存')
state(im,(430,620,740,230),'已保存待補處理','情緒或向量尚未齊備')
state(im,(430,1150,740,230),'衍生資料齊備','目前版本的情緒與向量已建立')
line(im,[(800,310),(800,620)],head='open');tag(im,845,382,'create [通過過濾]\n/ 保存文章',33)
line(im,[(800,850),(800,1150)],head='open');tag(im,845,928,'補評及向量建立\n[所需處理完成]',33)
line(im,[(430,1260),(180,1260),(180,730),(430,730)],head='open');tag(im,20,935,'update [文字變更]\n/ 清除舊向量\n依正文重設情緒',30)
line(im,[(1170,680),(1490,680),(1490,815),(1170,815)],head='open');tag(im,1200,490,'處理失敗\n/ 保留待處理\n供後續重試',31)
note(im,1430,'狀態由資料推得，非status欄位；只有留言變更時，不必然重設文章情緒。');save(im,'article_state')

mapping={'圖5-9 ':'analysis_layers','圖5-10 ':'analysis_collaboration','圖6-1 ':'trace','圖6-20 ':'orm','圖6-21 ':'crawler','圖6-22 ':'responsibility','圖7-1 ':'deployment','圖7-2 ':'packages','圖7-3 ':'components','圖7-4 ':'lock_state','圖7-5 ':'article_state'}
changed=[]
for i,p in enumerate(D.paragraphs):
 if p.style.name!='Figure Caption':continue
 for prefix,name in mapping.items():
  if not p.text.startswith(prefix):continue
  pic=D.paragraphs[i-1];assert pic._p.xpath('.//w:drawing'),p.text
  for r in list(pic.runs):r._r.getparent().remove(r._r)
  img=Image.open(F/(name+'.png'));ww=min(18,23*img.width/img.height)
  pic.add_run().add_picture(str(F/(name+'.png')),width=Cm(ww))
  pic.alignment=WD_ALIGN_PARAGRAPH.CENTER;pic.paragraph_format.first_line_indent=Pt(0);pic.paragraph_format.keep_with_next=True
  changed.append({'caption':p.text,'asset':name,'width_cm':ww})

# Clarify exactly what the redesigned notation represents.
for p in D.paragraphs:
 if p.text.startswith('類別方框分為名稱、屬性與操作區'):
  p.text='ORM類別圖節錄名稱與關鍵屬性；爬蟲類別圖節錄實際操作。空心三角箭頭指向父類別，實線關聯端標示多重性。圖中省略空白區隔以提高可讀性，完整欄位見第8章。PTTCrawler獨立實作，不畫成BrowserCrawler的子類別。'
 if p.text.startswith('套件以資料夾形狀表示'):
  p.text='套件以資料夾形狀表示，虛線箭頭由使用者指向被依賴者。圖中節錄主請求路徑：前端透過HTTP使用routers，routers依賴services，services依賴models與共用core模組。這不是完整匯入圖；core.scheduler另會協調crawlers與services，爬蟲也可能引用共用常數。不得將此摘要解讀為所有套件都嚴格單向分層。'
 if p.text.startswith('元件圖以«component»標示'):
  p.text='元件圖以«component»標示可辨識的功能單位，虛線開放箭頭由使用者指向被使用元件，旁註介面用途。HTTP API與WebSocket是不同通道，SQL介面不直接暴露给瀏覽器。圖中背景處理是應用內功能單位，不表示已部署為獨立服務；箭頭也不表示執行先後。'.replace('给','給')
 if p.text.startswith('本圖為UML產物依賴圖'):
  p.text='本圖為UML產物依賴圖。由上而下閱讀分析、設計、實作與驗證成果；«trace»虛線箭頭向上指向被追溯的產物，不是執行時資料流。各產物改動時，應核對其來源需求與後續實作、測試是否一致。'
# Place the responsibility table before its large diagram to avoid an isolated table page.
ps=D.paragraphs
fc=next(i for i,p in enumerate(ps) if p.style.name=='Figure Caption' and p.text.startswith('圖7-3 '))
anchor=ps[fc-1]._p
tc=next(p for p in ps if p.style.name=='Table Caption' and p.text.startswith('表7-2 '))
table=tc._p.getnext()
assert table.tag==qn('w:tbl')
anchor.addprevious(tc._p);anchor.addprevious(table)
assert len(changed)==11,changed
D.save(OUT)
(W/'polished-uml-manifest.json').write_text(json.dumps(changed,ensure_ascii=False,indent=2),encoding='utf8')
print('Updated',len(changed),'figures')
