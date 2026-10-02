from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import math,json,hashlib,re

W=Path(__file__).parent; ROOT=W.parent.parent
FIG=W/'chapter6-complete';FIG.mkdir(exist_ok=True)
D=Document(W.parent/'MeBOD_115年系統手冊_初審整合與程式節錄版.docx')
OUT=W.parent/'MeBOD_115年系統手冊_第6章UML與第9章程式完整版.docx'
FONT='C:/Windows/Fonts/kaiu.ttf'
def font(n):return ImageFont.truetype(FONT,n)
def lines(s,width,size):
 out=[]
 for line in s.split('\n'):
  cur=''
  for c in line:
   if font(size).getlength(cur+c)>width and cur:out.append(cur);cur=''
   cur+=c
  out.append(cur)
 return out
def label(d,x,y,s,width=400,size=32,center=False):
 ls=lines(s,width,size)
 for k,t in enumerate(ls):
  xx=x+(width-font(size).getlength(t))/2 if center else x
  d.text((xx,y+k*(size+7)),t,font=font(size),fill='#111827')
 return len(ls)*(size+7)
def dashed(d,a,b):
 length=math.dist(a,b)
 if length==0:return
 for pos in range(0,int(length),17):
  p=pos/length;q=min(pos+9,length)/length
  d.line((a[0]+(b[0]-a[0])*p,a[1]+(b[1]-a[1])*p,a[0]+(b[0]-a[0])*q,a[1]+(b[1]-a[1])*q),fill='#64748b',width=2)
def arrow(d,a,b,ret=False):
 if ret:dashed(d,a,b)
 else:d.line((a,b),fill='#334155',width=3)
 dx=b[0]-a[0];dy=b[1]-a[1];L=max(1,math.hypot(dx,dy));ux=dx/L;uy=dy/L
 pts=[(b[0]-19*ux-8*uy,b[1]-19*uy+8*ux),b,(b[0]-19*ux+8*uy,b[1]-19*uy-8*ux)]
 if ret:d.line(pts,fill='#334155',width=3)
 else:d.polygon(pts,fill='#334155')
def seq(name,actors,events,frames=()):
 width=2000;h=250+len(events)*135
 im=Image.new('RGB',(width,h),'white');dr=ImageDraw.Draw(im)
 xs=[170+i*1660/(len(actors)-1) for i in range(len(actors))]
 for x,t in zip(xs,actors):
  dr.rectangle((x-145,15,x+145,145),fill='#eff6ff',outline='#475569',width=2)
  label(dr,x-137,32,t,274,32,True);dashed(dr,(x,145),(x,h-20))
 for start,stop,kind,guards in frames:
  y=170+start*135;bottom=170+(stop+1)*135
  dr.rectangle((10,y,1990,bottom),outline='#64748b',width=2)
  dr.rectangle((10,y,115,y+36),fill='white',outline='#64748b',width=2)
  label(dr,23,y+2,kind,85,28)
  for row,guard in guards:
   gy=170+row*135
   if row!=start:dashed(dr,(10,gy),(1990,gy))
   label(dr,130,gy+3,guard,1810,29)
 for i,(a,b,t,ret) in enumerate(events):
  y=170+i*135
  assert a!=b
  lo=min(xs[a],xs[b]);hi=max(xs[a],xs[b]);maxw=hi-lo-34
  ls=lines(t,maxw,31)
  assert len(ls)<=2,(name,t,ls)
  label(dr,lo+17,y+40,t,maxw,31)
  if not ret:dr.rectangle((xs[b]-7,y+103,xs[b]+7,y+130),fill='#cbd5e1',outline='#64748b',width=2)
  arrow(dr,(xs[a],y+115),(xs[b],y+115),ret)
 im.save(FIG/(name+'.png'))

def method():
 im=Image.new('RGB',(2000,1180),'white');dr=ImageDraw.Draw(im)
 nodes=[(100,60,'«artifact»\n使用個案與需求\nUC01至UC06'),(1110,60,'«artifact»\n系統循序與事件合約\n輸入及後置條件'),(1110,435,'«artifact»\n詳細循序圖\n邊界 控制 實體'),(100,435,'«artifact»\n設計類別與責任\n屬性 操作 相依'),(100,820,'«artifact»\n原始碼與資料模型\n第8章與第9章'),(1110,820,'«artifact»\n驗證案例\n第10章')]
 for x,y,t in nodes:
  dr.rectangle((x,y,x+760,y+235),fill='#f8fafc',outline='#334155',width=3);label(dr,x+20,y+28,t,720,37,True)
 for a,b in [((1110,180),(860,180)),((1490,435),(1490,295)),((860,550),(1110,550)),((480,820),(480,670)),((1110,935),(860,935))]:arrow(dr,a,b,True)
 label(dr,80,1100,'虛線開放箭頭為«trace»依賴，箭頭指向被追溯的分析或設計產物。',1850,33)
 im.save(FIG/'method_trace.png')
method()

# Black-box contract diagrams show only external events and outcomes.
seq('contract_login',['使用者',':MeBOD'],[(0,1,'login(username, password)',False),(1,0,'Cookie與帳號資料',True),(1,0,'401或423；限流429',True)],[(1,2,'alt',[(1,'[帳號與密碼驗證成功]'),(2,'[驗證失敗或仍鎖定]')])])
seq('contract_query',['使用者或訪客',':MeBOD'],[(0,1,'searchDashboard(keyword, days, boards)',False),(1,0,'cached=true及儀表板資料',True),(1,0,'新計算結果；空集合亦可成功',True)],[(1,2,'alt',[(1,'[相同條件快取仍有效]'),(2,'[快取未命中]')])])
seq('contract_qa',['已登入使用者',':MeBOD'],[(0,1,'askQuestion(question, context, history)',False),(1,0,'回答及來源；用量已記錄',True),(1,0,'401／403／429或服務例外',True)],[(1,2,'alt',[(1,'[通過驗證且回答服務正常返回]'),(2,'[登入 額度 限流拒絕或例外外拋]')])])
seq('contract_crawl',['管理員',':MeBOD'],[(0,1,'startCrawl(platform, boards, pages)',False),(1,0,'started=true；背景工作已登記',True),(1,0,'400／403／409；不排入工作',True)],[(1,2,'alt',[(1,'[參數合法且取得租約]'),(2,'[參數 權限或租約不符合]')])])
seq('contract_personal',['已登入使用者',':MeBOD'],[(0,1,'getHistory(limit)',False),(1,0,'自己的歷史清單',True),(0,1,'deleteHistory(id)',False),(1,0,'刪除成功或404',True),(0,1,'addWatch(keyword, days)',False),(1,0,'本人監測詞或409／額度拒絕',True),(0,1,'exportArticles(filters)',False),(1,0,'方案允許範圍的XLSX',True),(0,1,'changePassword(old, new)',False),(1,0,'密碼更新；舊JWT失效',True)])
seq('session_check',['網頁邊界','auth服務\n«module»','User／DB'],[(0,1,'受保護請求攜帶Cookie',False),(1,2,'驗證JWT後查詢User',False),(2,1,'啟用狀態與密碼異動時間',True),(1,0,'回傳目前使用者，繼續端點',True),(1,0,'HTTP 401，要求重新登入',True)],[(3,4,'alt',[(3,'[JWT有效 帳號啟用 且未因改密碼失效]'),(4,'[權杖或帳號驗證失敗]')])])
seq('query_detail',['DashboardView','dashboard路由','快取／服務','Article／DB'],[(0,1,'GET /api/dashboard/full',False),(1,2,'正規化條件並查快取',False),(2,1,'命中：回傳既有資料',True),(1,2,'未命中：get_dashboard_full',False),(2,3,'同條件依序查統計及文章',False),(3,2,'查詢結果或空集合',True),(2,1,'結果並保存10分鐘快取',True),(1,0,'JSON資料與快取標示',True)],[(2,6,'alt',[(2,'[快取命中]'),(3,'[快取未命中]')])])
seq('history_detail',['分析／歷史畫面','analysis路由','分析服務','History／DB'],[(0,1,'GET /keyword（可選登入）',False),(1,2,'產生或取得分析摘要',False),(2,1,'result_json',True),(1,3,'登入者：新增user_id快照並commit',False),(1,0,'分析結果；訪客略過保存',True),(0,1,'GET /history（必須登入）',False),(1,3,'依目前user_id查詢與排序',False),(3,1,'本人歷史紀錄',True),(1,0,'序列化；部分指標即時計算',True)],[(3,3,'opt',[(3,'[已登入]')])])
seq('export_detail',['DashboardView','export路由','plan／export服務','Article／DB'],[(0,1,'GET /articles.xlsx及Cookie',False),(1,2,'檢查匯出權限及裁切天數',False),(2,1,'權限通過與有效範圍',True),(1,2,'正規化看板並取資料',False),(2,3,'依有效天數與平台查文章',False),(3,2,'符合條件的文章',True),(2,1,'建立XLSX位元組',True),(1,0,'StreamingResponse下載',True)])
seq('delete_password',['個人畫面','受保護路由','驗證服務','User／History DB'],[(0,1,'DELETE /history/{id}',False),(1,3,'以id且本人user_id讀取',False),(3,1,'符合的歷史或None',True),(1,3,'存在：刪除及稽核後commit',False),(1,0,'成功；不存在或非本人404',True),(0,1,'POST /auth/change-password',False),(1,2,'驗舊密碼及新密碼政策',False),(1,3,'雜湊 密碼異動時間及稽核commit',False),(1,0,'更新成功；後續舊JWT回401',True)])
seq('daily_detail',['排程觸發器','scheduler模組','租約／爬蟲','分析／預警'],[(0,1,'run_daily_job',False),(1,2,'啟用時try_acquire',False),(2,1,'取得租約；否則略過',True),(1,2,'逐看板續租 爬取及保存',False),(2,1,'爬取階段返回',True),(1,2,'finally依owner釋放租約',False),(1,3,'未停止時補評文章與留言',False),(1,3,'建立向量 再檢查預警',False),(3,1,'各階段結果',True),(1,0,'記成功日期並回摘要；關閉Session',True)])
seq('watch_detail',['MonitorView','monitor路由','plan服務','Watch／DB'],[(0,1,'POST /keywords（需登入）',False),(1,3,'同user_id與keyword查重',False),(3,1,'無重複；否則409',True),(1,2,'ensure_keyword_quota',False),(2,1,'尚有額度；否則拒絕',True),(1,3,'保存監測詞 user_id及稽核',False),(3,1,'commit後refresh',True),(1,0,'新監測詞JSON',True)])
seq('alert_detail',['排程／觸發端','alert_service','查詢／設定服務','Watch／Alert DB'],[(0,1,'run_alert_checks',False),(1,3,'讀啟用中的全部監測詞',False),(3,1,'watch清單',True),(1,2,'依關鍵字 天數查統計與門檻',False),(2,1,'文章數 負面比例及level',True),(1,3,'達門檻：按user_id及時間查重',False),(1,3,'無近期重複：新增Alert',False),(1,3,'有新增時commit及refresh',False),(1,0,'新預警清單',True)],[(3,6,'loop',[(3,'[逐一監測詞；不足門檻或重複時略過新增]')])])
seq('responsibility',['Vue邊界元件','FastAPI路由\n«module»','業務服務\n«module»','ORM／Session'],[(0,1,'使用者事件轉HTTP請求',False),(1,2,'驗證後委派業務操作',False),(2,3,'查詢或更新實體',False),(3,2,'實體及交易結果',True),(2,1,'業務資料或明確例外',True),(1,0,'HTTP狀態及JSON／檔案',True)])

def find(t):return next(p for p in D.paragraphs if p.text==t)
def add(a,t='',style='Normal'):return a.insert_paragraph_before(t,style)
diagram_manifest=[]
def figure(before,name,title,description):
 a=find(before)
 p=add(a);p.paragraph_format.keep_with_next=True;p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 im=Image.open(FIG/(name+'.png'));ratio=im.height/im.width
 width=min(18,19/ratio)
 p.add_run().add_picture(str(FIG/(name+'.png')),width=Cm(width))
 add(a,'圖6-X '+title,'Figure Caption')
 p=add(a,description);p.paragraph_format.keep_together=True
 diagram_manifest.append({'before':before,'image':name,'title':title})

figure('6-1-2 系統循序圖','method_trace','使用個案至程式與測試的追溯關係','本圖為UML產物依賴圖，不是執行時資料流。使用個案決定系統事件，詳細循序分配責任，設計類別與原始碼再以測試核對；各產物改動時沿追溯關係檢查影響。')
figure('6-1-3-2 查詢事件 searchDashboard','contract_login','login事件的成功與拒絕合約','此系統循序圖將MeBOD視為單一物件。前置條件、成功後置條件及拒絕結果與本節文字對應；內部帳號驗證互動另見6-1-4。')
figure('6-1-3-3 問答事件 askQuestion','contract_query','searchDashboard事件與快取結果','兩個alt區段互斥；即使沒有符合文章，仍可回傳合法空結果。此事件本身不寫入個人分析歷史。')
figure('6-1-3-4 爬取事件 startCrawl','contract_qa','askQuestion事件與用量後置條件','回答服務正常返回後才執行用量記錄；模型降級若仍以正常結果返回，同樣進入計次路徑。拒絕與外拋例外不能當成回答成功。')
figure('6-1-3-5 個人資料事件','contract_crawl','startCrawl事件的工作接受邊界','started=true表示已接受背景工作；文章保存、情緒評分及向量建立在後續執行，不能把此回應解讀為全部完成。')
figure('6-1-4 登入與工作階段詳細流程','contract_personal','個人歷史監測匯出及改密碼事件','圖中五組請求與回應為各自獨立的操作示例，並非必須連續完成的單一交易。每次請求都重新驗證帳號；各事件的內部流程分別見6-1-5及6-1-8。')
figure('6-1-5 查詢 分析 歷史與匯出的完整流程','session_check','受保護請求的工作階段驗證','對應auth_service.get_current_user。每次受保護操作都驗證權杖與User狀態；改密碼後的舊權杖即使未到期，也因password_changed_at檢查而被拒絕。')
for name,title,desc in [
 ('query_detail','儀表板篩選快取與文章聚合循序','對應dashboard.py與dashboard_service.py。查詢條件包含平台看板、關鍵字、天數及排序；聚合查詢採同步Session依序執行，不表示各指標平行查詢。'),
 ('history_detail','分析保存與個人歷史讀取循序','對應analysis.py。opt片段僅在登入時建立快照；讀取歷史需登入且限定本人。圖中摘要JSON保存與後續查詢為不同請求，部分歷史指標會依現況重算。'),
 ('export_detail','Excel權限與有效查詢範圍循序','對應export.py、plan_service.py及export_service.py。未登入或無匯出權限時提前拒絕，不進入建檔流程；檔名days可能仍為原請求值，內容則依裁切後天數。'),
 ('delete_password','刪除個人歷史與變更密碼循序','兩組請求各自獨立。不存在或非本人的歷史一律回404；改密碼必須先驗舊密碼及新政策，再提交雜湊與異動時間。錯誤分支提前結束，不執行後續寫入。'),
]:figure('6-1-6 問答與混合檢索的詳細流程',name,title,desc)
figure('6-1-8 監測與預警流程及現況限制','daily_detail','每日排程的租約範圍與後處理','對應scheduler.run_daily_job。停用或取鎖失敗會提前略過；停止要求會在爬取後提前返回。每日流程在後處理前放鎖，與手動工作的租約涵蓋範圍不同；最外層finally負責關閉Session。')
figure('6-2 設計類別圖與責任分配','watch_detail','本人監測詞建立與方案檢查','對應monitor.add_keyword。去空白與請求驗證後先查本人重複，再檢查方案組數；成功保存user_id、監測設定與稽核紀錄。此圖只描述新增操作，不代表其他預警異動已完成同樣隔離。')
figure('6-2 設計類別圖與責任分配','alert_detail','風險評估與帳號別預警去重','對應alert_service.py。文章數先達最低要求，再依critical優先於warning判定；同帳號、關鍵字與時間窗防止重複新增。GET預警清單另依user_id過濾，手動全域檢查及已讀操作的現況限制仍依本節說明。')
figure('第7章 實作模型','responsibility','設計責任如何落實到邊界控制服務與實體','本循序圖補充前述設計類別圖的協作方式。路由與服務以«module»標示函式模組，沒有虛構同名Python類別；ORM與Session承接資料責任，實際交易由呼叫端決定，不能把所有請求都畫成自動commit。')

# Number all chapter-six captions in reading order; Word refreshes caption TOCs.
num=0
for p in D.paragraphs:
 if p.style.name=='Figure Caption' and p.text.startswith('圖6-'):
  num+=1;t=re.sub(r'^圖6-[\dX]+',f'圖6-{num}',p.text)
  if p.runs:
   p.runs[0].text=t
   for r in p.runs[1:]:r.text=''
  else:p.add_run(t)

code_manifest=[]
def snippet(before,title,path,first,last,explanation):
 a=find(before);src=(ROOT/path).read_text(encoding='utf-8').splitlines()
 start=next(i for i,s in enumerate(src) if first in s)
 if 'created = embed_pending_articles(' in first:start-=1
 stop=next(i for i in range(start,len(src)) if last in src[i])
 code='\n'.join(src[start:stop+1])
 p=add(a,title);p.paragraph_format.keep_with_next=True
 for r in p.runs:r.bold=True
 p=add(a,f'原碼節錄：{path}，第{start+1}至{stop+1}行。僅節錄本項責任的核心程式，其餘匯入與流程見原檔。');p.paragraph_format.keep_with_next=True
 p=add(a,code);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=1
 for r in p.runs:
  r.font.name='Consolas';r.font.size=Pt(10);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'標楷體')
 p=add(a,explanation);p.paragraph_format.keep_together=True
 code_manifest.append({'section_before':before,'title':title,'path':path,'start':start+1,'end':stop+1,'sha256':hashlib.sha256(code.encode()).hexdigest()})

snippet('9-1-2 儀表板、文章與個人歷史元件','程式9-1-1 身分與密碼異動檢查','backend/app/services/auth_service.py','    payload = decode_access_token(token)','    return user',
 '本段位於get_current_user，先解碼JWT，再查User並檢查啟用與密碼異動時間，成功才回傳帳號物件。失敗回401；前置的_read_token已處理Cookie或Bearer來源。對應6-1-4工作階段圖。測試應涵蓋無效簽章、停用帳號及改密碼後舊JWT失效。')
snippet('9-1-3 AI問答與檢索元件','程式9-1-2A 儀表板條件與快取命中','backend/app/routers/dashboard.py','    selected_boards = normalize_filter_boards(boards)','        }',
 '四種篩選資訊一起形成快取鍵；命中可直接回傳，未命中才呼叫get_dashboard_full並保存10分鐘。此公開端點不建立個人歷史，對應6-1-5查詢循序。')
snippet('9-1-3 AI問答與檢索元件','程式9-1-2B 已登入者保存分析快照','backend/app/routers/analysis.py','    if current_user is not None:','        db.commit()',
 '此段屬analyze_keyword，不是dashboard_full。訪客跳過整段；登入者將user_id、查詢條件與result_json存入AnalysisHistory。驗證訪客不新增紀錄、兩個帳號互不讀到對方紀錄；讀取隔離的節錄另見9-1-8。')
snippet('9-1-4 爬取、文章更新與租約元件','程式9-1-3 回答快取與Dashboard脈絡分支','backend/app/services/rag_service.py','    cache_key = _qa_cache_key(question, dashboard_context, history)','        answer = generate_dashboard_context_answer(question, dashboard_context, history)',
 '本段位於answer_question：use_cache為真且快取存在便直接返回；未命中時，有dashboard_context才進入脈絡回答。沒有脈絡的else分支另走意圖、關鍵字與向量混合檢索。即使快取命中，路由仍會記用量。對應6-1-6分支圖及9-1-7控制流程。')
snippet('9-1-5 匯出、方案、監測與管理元件','程式9-1-4A 過期租約的原子接手','backend/app/services/lock_service.py','    acquired = db.query(SystemLock).filter(','    return bool(acquired)',
 '前面的唯一鍵插入失敗且rollback後才執行本段。UPDATE同時限制name與expires_at，藉受影響列數判斷是否取得；成功後記錄本次UUID owner。此設計需以PostgreSQL並行測試驗證，不能由SQLite單一請求推論。')
snippet('9-1-5 匯出、方案、監測與管理元件','程式9-1-4B 釋放時核對租約擁有者','backend/app/services/lock_service.py','def release(db: Session','        db.info["lock_owners"].pop(name, None)',
 '一般工作只刪除自己owner對應的鎖，避免舊工作結束時刪掉已被新工作接手的租約。沒有owner則直接返回。管理員force_release是另一條維運路徑，不能與此一般釋放混用。對應6-1-7及第7章租約狀態機。')
snippet('9-1-6 問答請求資料與驗證節錄','程式9-1-5A 匯出權限與方案天數','backend/app/routers/export.py','    plan_service.ensure_export_allowed(db, current_user)','    filename = f"articles_{days}d.xlsx"',
 '匯出先檢查權限，再將clamp_history_days回傳值用於資料查詢；畫面天數與檔名可能仍為原請求值。測試必須核對XLSX內的日期及平台，而非只檢查下載成功。對應6-1-5匯出循序。')
snippet('9-1-6 問答請求資料與驗證節錄','程式9-1-5B 監測詞歸屬與額度','backend/app/routers/monitor.py','    plan_service.ensure_keyword_quota(db, current_user)','    db.refresh(watch)',
 '同名檢查已在前段完成；此處通過額度後保存本人user_id並寫入稽核。對應6-1-8新增監測詞循序。監測詞與預警依關鍵字等業務資訊協作，不能據此宣稱Alert具有watch_keyword_id外鍵。')
snippet('9-1-6 問答請求資料與驗證節錄','程式9-1-5C 管理員角色防線','backend/app/services/auth_service.py','def require_admin(current_user: User','    return current_user',
 'require_admin先依get_current_user取得可信身分，再拒絕非admin角色。畫面隱藏按鈕不能代替此檢查；應驗證一般帳號直接呼叫管理端點時仍回403。方案額度與管理員角色是分開的判斷。')
snippet('9-2-2 資料庫、遷移與啟停元件','程式9-2-1 共用API連線','frontend/src/services/api.js','const apiBaseUrl =','});',
 '所有使用此api物件的頁面共用baseURL、逾時與Cookie設定；跨來源時仍需伺服器允許相應憑證政策。HttpOnly只限制JavaScript讀取Cookie，不代表XSS無法代使用者送出請求。401回應攔截見9-2-4，對應6-2責任協作圖。')
snippet('9-2-3 補建腳本、測試與建置元件','程式9-2-2A 請求Session生命週期','backend/app/core/database.py','    db = SessionLocal()','        db.close()',
 'get_db以yield提供一次請求的Session，finally保證關閉；本段沒有自動commit，交易由路由或服務明確提交。資料庫例外與業務例外均不得跳過資源釋放。對應第7章啟動請求關閉程序。')
snippet('9-2-3 補建腳本、測試與建置元件','程式9-2-2B 遷移新增最近爬取時間','backend/alembic/versions/f1a92b3c4d56_add_article_last_crawled_at.py','def upgrade() -> None:','    # Old rows stay NULL:',
 '升級先讀現有欄位，只在缺欄位時新增nullable DateTime。舊資料保留NULL，避免捏造既有文章最後抓取時間；版本鏈由Alembic管理。應驗證空庫及舊庫升級後欄位存在且既有內容保留。')
snippet('9-2-4 共同錯誤合約與修改程序','程式9-2-3A 批次向量補建與重試','backend/scripts/backfill_embeddings.py','                created = embed_pending_articles(','            remaining = count_pending_embeddings(db)',
 '此段在backfill迴圈內執行一批向量生成；例外時記錄日誌、等待後重試，再檢查待處理篇數。外層另有limit與停滯偵測，但本段例外分支直接continue，不能宣稱所有持續失敗都會在三次後停止。維運先dry-run再依額度執行。')
snippet('9-2-4 共同錯誤合約與修改程序','程式9-2-3B PostgreSQL競爭測試核心','backend/system_tests/test_postgres.py','    barrier = Barrier(4)','        assert list(workers.map(acquire, range(4))).count(True) == 1',
 '四個執行緒各用獨立Session，Barrier讓取得租約的時機接近，斷言僅一個True。外層parametrize另涵蓋空鎖與過期鎖。此為既有測試程式節錄，不代表本次文件修訂重新執行測試。對應第10章PG06及PG07。')
snippet('第10章 測試模型','程式9-2-4A 共用401回應處理','frontend/src/services/api.js','api.interceptors.response.use(','export default api;',
 '只有本機曾保存auth_user且收到401時才清除顯示狀態並導向登入；訪客401留給各頁面提示。localStorage僅用於畫面判斷，不能作授權憑證。Promise.reject保留錯誤供呼叫端處理，對應9-1-9訪客頁面防護。')
snippet('第10章 測試模型','程式9-2-4B 資源不存在的HTTP合約','backend/app/routers/articles.py','    article = db.query(Article).filter(','        raise HTTPException(status_code=404, detail="找不到此文章。")',
 '文章不存在時回404，而非以200及空物件掩蓋。修改此契約時需同步檢查ArticleDetailView與測試。其他401、403、409、422及429分別代表身分、權限、衝突、輸入與頻率限制，前端應依意義呈現。')

# Keep added explanations and original snippets distinct; no application changes.
add(find('9-1 元件清單及其規格描述'),'本版於9-1-1至9-1-5及9-2-1至9-2-4各自附對應原碼。程式編號按所在小節排列，便於與第6章UML互相核對；9-1-6至9-1-10保留為特定實作的延伸節錄。原碼以等寬字呈現，與說明正文分開。')
for p in D.paragraphs:
 if p.text.startswith('程式與證據：第9章補入五段'):
  p.text='程式與證據：第9章各元件及附屬元件均附核心原碼與流程解說，不刊載完整專案。第6章各流程以相應UML補充。初審文件不是評審意見表，不能改寫成教師已提出的意見。'
for p in D.paragraphs:
 if p.text.startswith('原碼節錄：'):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 if p.text.startswith('9-1-7 '):p.paragraph_format.page_break_before=False
D.save(OUT)
(W/'chapter6-9-manifest.json').write_text(json.dumps({'diagrams':diagram_manifest,'new_excerpts':code_manifest,'chapter6_figures':num},ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved document; new diagrams',len(diagram_manifest),'new snippets',len(code_manifest),'chapter6 figures',num)

