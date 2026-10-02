from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import json,hashlib
W=Path(__file__).parent;ROOT=W.parent.parent
D=Document(W/'rag-base.docx');a=next(p for p in D.paragraphs if p.text=='9-2 其他附屬之各種元件')
def para(t,style='Normal'):
 return a.insert_paragraph_before(t,style)
p=para('9-1-11 RAG檢索增強生成的核心實作','Heading 3');p.paragraph_format.page_break_before=True
para('本節補充9-1-3的問答元件與9-1-10的RRF融合，節錄目前專案的RAG核心路徑。RAG是先從資料庫找出相關文章，再把文章證據交給生成模型回答。執行順序為意圖解析、硬條件篩選、關鍵字與向量檢索、RRF排名融合、文章壓縮、模型生成及來源回傳。以下皆為原檔連續節錄，不刊載完整模組。')
para('適用分支：answer_question未命中可用快取，且沒有dashboard_context時才執行此RAG路徑；若有儀表板脈絡，則直接依脈絡生成。身分、額度及用量處理由qa路由負責，見9-1-7。對應第6章圖6-15。')
manifest=[]
def code(title,path,first,last,desc):
 src=(ROOT/path).read_text(encoding='utf8').splitlines();start=next(i for i,s in enumerate(src) if first in s);end=next(i for i in range(start,len(src)) if last in src[i]);t='\n'.join(src[start:end+1])
 p=para(title);p.paragraph_format.keep_with_next=True
 for r in p.runs:r.bold=True
 p=para(f'原碼節錄：{path}，第{start+1}至{end+1}行。');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.keep_with_next=True
 p=para(t);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 for r in p.runs:r.font.name='Consolas';r.font.size=Pt(10);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'標楷體')
 p=para(desc);p.paragraph_format.keep_together=True
 manifest.append(dict(title=title,path=path,start=start+1,end=end+1,sha256=hashlib.sha256(t.encode()).hexdigest()))
R='backend/app/services/rag_service.py'
code('程式9-1-11A RAG主路徑',R,'        intent = parse_question_intent(question)','            "cached": False,','先解析question得到intent，retrieve_articles依意圖找證據，再交generate_rag_answer生成回答。此節錄位於else分支，末尾回傳字典的閉合及快取保存見原檔。sources由serialize_sources產生，預設最多列5篇；檢索預設最多12篇，因此來源清單是節錄，不等同完整模型輸入。confidence是模型或降級結果的標示，不是經校準的正確率。')
code('程式9-1-11B 混合檢索與降級',R,'    keyword_results = retrieve_by_keyword(db, intent, limit)','    return fuse_rankings([keyword_results, vector_results], limit)','關鍵字與向量檢索使用同一組意圖條件。沒有question或向量結果為空時，回傳關鍵字結果；兩者皆有結果時用RRF融合。關鍵字端在標題與內文比對詞彙，先依命中分數再依熱度或時間排序；向量端處理同義與不同措辭。RRF實作見9-1-10：同article.id合併，名次由1起算，常數RRF_K=60，不能把詞彙分數與餘弦相似度直接相加。')
code('程式9-1-11C 語意候選與排序還原',R,'    candidate_ids = [','    ordered = [by_id[article_id] for article_id in ranked_ids if article_id in by_id]','本段位於retrieve_by_vector，先以apply_hard_filters限制候選文章，再計算語意相似度。SQL IN查詢不保證順序，因此依ranked_ids重建排序。後續迴圈將命中段落附在暫時性matched_chunk屬性，供文章脈絡整理使用，不會新增資料庫欄位。')
code('程式9-1-11D 餘弦相似度與文章去重','backend/app/services/embedding_service.py','    matrix = np.frombuffer','    return ranked[:limit]','本段位於rank_by_similarity：ArticleChunk的float32位元組還原為矩陣，正規化後與問題向量計算餘弦相似度，1e-12避免零除。同篇文章有多個段落時只保留最高分及其原文，再取前limit篇。這是NumPy計算，不是pgvector索引；問題向量生成失敗時，前段會回空清單讓上層退回關鍵字檢索。')
code('程式9-1-11E 建立有長度限制的文章脈絡',R,'    articles_context = compress_articles_for_llm(','        max_comment_chars=300,','本段的函式呼叫閉合位於下一行。每篇主文與留言各以300字為上限，整體文章脈絡預算為13000字元；這些是字元限制，不等於token數。後續prompt加入問題、解析意圖、壓縮文章及最近對話，要求繁體中文JSON與依文章證據回答。若無相關文章，函式在前段直接返回資料不足與low信心，不呼叫生成模型。')
code('程式9-1-11F 生成與降級回覆',R,'    fallback = _fallback_answer(question, intent, articles)','        return fallback','generate_json_response將組好的prompt交給模型；_safe_json解析結果，JSON不合法時採fallback，呼叫外拋例外亦回fallback。這種降級維持介面可用，但不代表產生了完整模型分析；現有正常返回仍由路由記錄用量。提示詞要求依證據回答，也不能保證模型絕不產生錯誤，使用者仍須核對來源。')
para('驗證重點：①相同硬條件下比較詞彙、向量與融合結果；②向量不存在或問題向量失敗時仍可退回詞彙；③同篇多段不重複回傳且保存最佳段落；④無文章時不生成推測答案；⑤模型例外與不合法JSON有降級結果；⑥有dashboard_context或有效快取時不誤走此路徑。上述為程式解讀及應驗證項目，本次文件增補未重新執行測試。')
# Include closing lines for excerpts without changing quoted source.
ps=D.paragraphs
for i,p in enumerate(ps):
 if p.text.startswith('程式9-1-11A') or p.text.startswith('程式9-1-11E'):
  prov=ps[i+1];cp=ps[i+2];e=next(e for e in manifest if e['title']==p.text);src=(ROOT/e['path']).read_text(encoding='utf8').splitlines();e['end']+=1;cp.runs[0].text='\n'.join(src[e['start']-1:e['end']]);e['sha256']=hashlib.sha256(cp.text.encode()).hexdigest();prov.text=f"原碼節錄：{e['path']}，第{e['start']}至{e['end']}行。"
for p in D.paragraphs:
 if '末尾回傳字典的閉合及快取保存見原檔。' in p.text:p.text=p.text.replace('末尾回傳字典的閉合及快取保存見原檔。','後續快取保存見原檔。')
 if '本段的函式呼叫閉合位於下一行。' in p.text:p.text=p.text.replace('本段的函式呼叫閉合位於下一行。','')
D.save(W.parent/'MeBOD_115年系統手冊_RAG程式增補版.docx')
(W/'rag-excerpts-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print('Added six RAG excerpts')
