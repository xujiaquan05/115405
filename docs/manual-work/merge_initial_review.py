from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import hashlib, json

W=Path(__file__).parent
ROOT=W.parent.parent
D=Document(W.parent/'MeBOD_115年系統手冊_網站操作與A4修訂版.docx')
OUT=W.parent/'MeBOD_115年系統手冊_初審整合與程式節錄版.docx'
def find(t):return next(p for p in D.paragraphs if p.text==t)
def add(a,t,style='Normal'):
 p=a.insert_paragraph_before(t,style);return p
def replace(p,t):
 if p.runs:
  p.runs[0].text=t
  for r in p.runs[1:]:r.text=''
 else:p.add_run(t)

for p in D.paragraphs:
 t=p.text
 if t.startswith('指導老師：'):replace(p,'指導老師：陳信宏 老師')
 elif t.startswith('組　　長：'):replace(p,'組　　長：11246044 陳威帆')
 elif t.startswith('組　　員：'):replace(p,'組　　員：11246039 許家全')
 elif t.strip()=='＿＿＿＿＿＿＿＿':replace(p,'')
 elif t.startswith('本专題需要') or t.startswith('本專題需要專案整合'):
  replace(p,'本組為雙人團隊，組長為11246044陳威帆，組員為11246039許家全，指導老師為陳信宏老師。以下角色表說明目前系統的交付責任；具名工作與比例承接民國115年6月2日初審手冊，屬初審階段紀錄。')
 elif t.startswith('具名分工表須由團隊'):
  replace(p,'初審分工記載許家全負責程式架構、前後端與AI問答整合，陳威帆負責資料蒐集、文件與報告。初審原文另列「模型訓練」，目前系統實作為呼叫Gemini既有模型，不據此宣稱完成自行訓練；最終分工應依後續實際產出確認。')
 elif t.startswith('本章依「9-1'):
  replace(p,'本章依「9-1 元件清單及其規格描述」及「9-2 其他附屬之各種元件」組織，先說明責任、介面、相依性及例外，再以少量專案原始碼節錄說明實作。節錄附檔案與行號，省略處均明確標示；完整程式仍以專案交付，不在手冊逐檔刊載。〔AI02〕')
 elif t.startswith('人工智慧使用工具、範圍、頁碼統一') or t.startswith('人工智慧使用工具、範圍、頁碼及內文序號統一'):
  replace(p,'人工智慧使用工具、範圍與頁碼列於第14章。初審手冊已揭露ChatGPT與Claude的使用，本版保留為初審階段紀錄，並與本次Codex修訂及系統執行時的Gemini分別標示；初審頁碼不直接沿用至本版。')
 elif t.startswith('以下資料仍須由團隊補齊'):
  replace(p,'初審名冊與當時貢獻比例已整合。最終分工與比例仍須由團隊確認；各組員GitHub截圖、正式評審意見及最終繳交附件應依實際紀錄補齊。')
 elif t.startswith('提交前確認封面的教師'):
  replace(p,'提交前確認初審名冊是否異動、最終分工與貢獻比例，以及評審修正紀錄。資料庫格式若仍要求MDF及LDF，應與指導單位確認PostgreSQL備份的替代方式。')

for tb in D.tables:
 if tb.cell(0,0).text=='身分及學號姓名':
  vals=[['身分及學號姓名','初審記載工作內容','初審貢獻度'],['組長 11246044 陳威帆','資料蒐集、文件製作、報告製作；原列模型訓練之界定見下文。','40%'],['組員 11246039 許家全','主程式架構、使用者介面、前後端開發、AI串接與問答系統。','60%'],['合計','民國115年6月2日初審紀錄；最終比例待團隊確認。','100%']]
  while len(tb.rows)>4:tb._tbl.remove(tb.rows[-1]._tr)
  for row,vs in zip(tb.rows,vals):
   for cell,value in zip(row.cells,vs):replace(cell.paragraphs[0],value)
 for row in tb.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    if p.text=='具備；姓名待填':replace(p,'具備；名冊依初審補入')
    elif p.text=='內容具備；具名紀錄待填':replace(p,'具備；初審分工已整合')

a=find('第2章 營運計畫')
add(a,'1-5 初審成果與目前版本','Heading 2')
add(a,'初審以PTT文章蒐集、Dashboard、AI問答、History及Excel匯出作為核心成果。目前版本沿用同一研究目的，擴充Dcard、Mobile01與Threads爬蟲、登入與角色管理、個人監測預警、文章留言及向量切段。四平台代表程式支援範圍，實際收錄狀況仍需查閱工作紀錄。')
add(a,'使用情境以診所或品牌分析人員為主：先設定療程或品牌關鍵字與觀察期間，閱讀趨勢及情緒，再開啟代表文章核對來源；需要追蹤時建立監測詞，需對外報告時匯出或列印。AI答案是整理線索，來源文章及查詢範圍才是結論回查的依據。')
add(a,'初審提及OpView與KEYPO作為同類產品。本專題比較重點為醫美領域工作流程、部署可控性與來源回查，不沿用未驗證的競品價格或效能優劣。付費方案與營運成本仍為後續驗證事項。')
a=find('第2章 營運計畫')
add(a,'社會價值：彙整公開心得可協助使用者辨認討論焦點與負面訊號。文章聲量不等於療程安全性或醫療成效；報告應附期間、平台與原文來源，避免將模型生成文字當成醫療建議。此定位承接初審的消費資訊透明目標。')
a=find('5-2 使用個案圖')
add(a,'初審曾以一般API兩秒、AI問答十五秒及每月99%可用性作為目標。這些數字保留為候選驗收目標，不標示為已達成結果。效能測試須固定資料量、請求組合與併發數，分開量測快取命中及未命中，記錄P50、P95與錯誤率；可用性另依持續監測期間計算。')

manifest=[]
def excerpt(anchor,title,path,start,end,explain,test):
 p=add(anchor,title,'Heading 3');p.paragraph_format.page_break_before=not title.startswith('9-1-6 ')
 add(anchor,explain)
 add(anchor,f'原始碼位置：{path}，第{start}至{end}行；民國115年9月30日工作目錄版本。下列為連續節錄，函式其餘部分與匯入宣告請參閱原檔。')
 lines=(ROOT/path).read_text(encoding='utf-8').splitlines()[start-1:end]
 code=add(anchor,'')
 code.paragraph_format.first_line_indent=Pt(0)
 code.paragraph_format.space_after=Pt(8)
 code.paragraph_format.line_spacing=1
 code.paragraph_format.keep_together=True
 for i,line in enumerate(lines):
  r=code.add_run(line+('\n' if i<len(lines)-1 else ''));r.font.name='Consolas';r.font.size=Pt(10)
  r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'標楷體')
 add(anchor,'流程與驗證：'+test)
 manifest.append({'path':path,'start':start,'end':end,'sha256':hashlib.sha256(('\n'.join(lines)).encode()).hexdigest()})

a=find('9-2 其他附屬之各種元件')
excerpt(a,'9-1-6 問答請求資料與驗證節錄','backend/app/routers/qa.py',27,38,
 'QuestionRequest屬控制層輸入契約，將問題、Dashboard脈絡與歷史訊息交給路由使用。Pydantic欄位約束先限制格式與長度，避免服務層接收任意角色或過長訊息。',
 'HistoryItem只允許user及assistant；問題長度為2至500字，歷史單筆上限2000字。錯誤型別或長度由請求驗證拒絕；此處不代表登入、額度或語意品質已通過，後續仍由路由檢查。對應UC03與第10章問答案例。')
excerpt(a,'9-1-7 問答控制與額度順序節錄','backend/app/routers/qa.py',41,66,
 'ask_question以相依注入取得Session及登入帳號，先檢查方案額度，再呼叫回答服務。no_cache轉換成use_cache，讓重新生成請求可略過回答快取。',
 '驗證順序為限流、登入、方案額度、回答與用量記錄。原碼註解寫成功才計次，實際界線是answer_question正常返回；快取答案及正常返回的降級答案仍可能計次，不能推論所有模型失敗都不扣額度。')
excerpt(a,'9-1-8 個人歷史資料隔離節錄','backend/app/routers/analysis.py',298,308,
 'analysis_history使用AnalysisHistory.user_id限定資料擁有者，再依時間倒序及limit取回紀錄。此段位於已由get_current_user保護的路由內，說明控制層授權如何落實到資料查詢。',
 '甲帳號建立分析後，乙帳號查詢不應取得甲的紀錄；未登入回401。不要把user_id改成前端可任意指定的值，也不要直接改查共用analysis_results。對應UC05、第8章analysis_history及第10章隔離案例。')
excerpt(a,'9-1-9 訪客歷史畫面防護節錄','frontend/src/views/HistoryView.vue',300,315,
 'HistoryView的fetchHistory先檢查isAuthenticated。訪客清空本頁紀錄並顯示登入提示，直接return，因此不會向私人歷史端點送出原先造成401提示的請求。',
 '以訪客開啟History時，應顯示登入訊息且不呼叫GET /api/analysis/history；已登入者才進入讀取流程。前端防護改善操作體驗，後端仍必須保留身分驗證及資料隔離。')
excerpt(a,'9-1-10 混合檢索排名融合節錄','backend/app/services/rag_service.py',423,433,
 'fuse_rankings將關鍵字與向量檢索結果依文章id合併，以RRF名次倒數加總排序，避免將不同尺度的原始分數直接相加。此段為排名融合核心，RRF_K在模組中設定為60。',
 '每組排名從1開始，文章在多組結果出現時累加分數；by_id保存對應Article物件，最後只取limit筆。驗證重複文章不重複輸出、空排名回空清單、結果數不超過limit。此節對應UC03未提供Dashboard脈絡的檢索分支。')

a=find('14-1 人工智慧輔助使用說明')
add(a,'[10] 第115405組。醫美時尚輿情分析系統系統手冊初審版。民國115年6月2日。用於團隊名冊、初審分工、專題動機及既有AI揭露之版本對照。')
add(a,'初審AI使用紀錄〔AI03〕：ChatGPT協助PTT解析程式骨架與PostgreSQL連線除錯；Claude協助手冊文字、UML章節及操作步驟整理。以上為初審文件已載明的歷史使用範圍，並非本次新增工具操作紀錄。')
a=next(p for p in D.paragraphs if p.text.startswith('附錄') and 'C' in p.text and '交付與維護檢核' in p.text and p.style.name.startswith('Heading'))
add(a,'初審與本版差異核對','Heading 3')
for t in [
 '名冊與分工：補入陳信宏老師、陳威帆與許家全；保留初審40%與60%比例的日期及範圍，最終貢獻不由提交次數推算。',
 '權限與歷史：初審的一般使用者、行銷人員與管理員敘述改以實際訪客、帳號與管理員授權為準。爬蟲限管理員；AI對話與個人分析歷史為不同功能，不宣稱所有問答都寫入History頁。',
 '實作與資料庫：不沿用初審假設的QASession、CrawlJob或取消狀態作為實際ORM。資料庫以目前模型與Alembic為準，保留第8章十九表總覽及十四條實體外鍵。',
 '部署與版本：操作畫面採目前mebod.clouda.dpdns.org網站；初審Render網址與舊114儲存庫僅屬歷史資料，不能替代目前安裝設定。',
 '程式與證據：第9章補入五段專案原碼及流程解說，不刊載完整模組。初審文件不是評審意見表，未將文件內容改寫成教師已提出的意見。',
]:add(a,t)
D.save(OUT)
(W/'initial-merge-code-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('Document saved')


