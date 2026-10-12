"""第六輪：依程式與 Git 實際紀錄校正事實。

- 甘特圖右側的月份里程碑改為各功能第一次提交的月份（git log --diff-filter=A）
- GitHub 提交數更新（含本輪提交）
- 備忘稿移除已刪頁面與舊時間的描述

輸入 stage5c.pptx，輸出 stage6.pptx。
"""
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)


def set_text(sh, text):
    par = sh.text_frame.paragraphs[0]
    if not par.runs:
        par.add_run()
    par.runs[0].text = text
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)


def by_text(slide, text):
    return next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip() == text)


# 甘特圖里程碑：依各檔案首次提交日期
# ptt_crawler 05-12、websocket 05-20、auth_service 07-12、dcard 08-09、mobile01 08-19、
# threads 08-20、plan_service 09-03、alembic 09-08、embedding_service 09-21
s = p.slides[15]
for old, new in [
    ("PTT 爬蟲、儀表板雛形", "PTT 爬蟲、儀表板、WebSocket"),
    ("帳號、WebSocket、Threads", "帳號登入與角色權限"),
    ("Dcard、Mobile01、方案額度", "Dcard、Mobile01、Threads"),
    ("Alembic 遷移、向量檢索", "方案額度、Alembic、向量檢索"),
    ("整合測試、情緒一致性修正", "情緒一致性、預警權限修正"),
]:
    set_text(by_text(s, old), new)

MONTHS = [14, 0, 27, 42, 47, 34]  # 10 月含本輪提交
chart = next(sh for sh in s.shapes if sh.has_chart).chart
cd = CategoryChartData()
cd.categories = ["5月", "6月", "7月", "8月", "9月", "10月"]
cd.add_series("提交數", MONTHS)
chart.replace_data(cd)
tf = chart.chart_title.text_frame
tf.paragraphs[0].runs[0].text = f"GitHub 每月提交（共 {sum(MONTHS)} 次）"
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("GitHub："):
        set_text(sh, f"GitHub：7–10 月 {sum(MONTHS[2:])} 次提交（10/12 統計，皆已推送）；非逐人貢獻比例。")
s.notes_slide.notes_text_frame.paragraphs[0].runs[0].text = (
    "【陳威帆｜45 秒】甘特圖保留原稿115/2至116/1的計畫，10月後為計畫工作。繳件與展覽日期沿用原稿，發表前仍按學校公告確認。"
    "右側里程碑依各功能第一次提交的月份：5月PTT爬蟲、儀表板與WebSocket；7月帳號登入；8月Dcard、Mobile01、Threads；"
    "9月方案額度、Alembic與向量檢索；10月情緒一致性與預警權限修正。"
    f"提交數由GitHub於10/12統計：5月14、6月0、7月27、8月42、9月47、10月34，共{sum(MONTHS)}次，皆已推送；"
    "提交集中於單一帳號，不能當作兩位組員的貢獻比例。")

# 備忘稿：第 2 頁與結論頁的時間、已刪除的測試頁
n2 = p.slides[1].notes_slide.notes_text_frame.paragraphs[0].runs[0]
n2.text = "【陳威帆｜15 秒】先談問題與營運，再由家全說明技術、功能與Demo，最後回到進度與限制；附錄供問答使用。"
n18 = p.slides[17].notes_slide.notes_text_frame.paragraphs[0].runs[0]
n18.text = n18.text.replace("RAG與分數規則已可說明，測試報告可核對", "RAG與分數規則已可說明，後端493項、前端43項自動化測試通過").replace(
    "結束於14:30，保留30秒緩衝。", "原第16頁測試頁已刪除，主講約13分55秒，可在Demo或此頁多說約35秒。")

p.save(OUT)
print("saved", OUT, sum(MONTHS), sum(MONTHS[2:]))
