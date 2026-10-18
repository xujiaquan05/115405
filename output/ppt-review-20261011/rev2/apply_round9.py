"""第九輪：第 7 頁改為「成本明細表＋計算式＋收支對照表」，讓成本來源一眼看懂。

輸入 stage8.pptx，輸出 stage9.pptx。數字與 apply_round8.py 相同（Gemini 3.1 Flash-Lite）。
"""
import math
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)
s = p.slides[6]

# ── 數字（與第八輪一致） ─────────────────────────────
RATE, IN_PRICE, OUT_PRICE = 32, 0.25, 1.50


def cost(tokens_in, tokens_out):
    return (tokens_in * IN_PRICE + tokens_out * OUT_PRICE) / 1e6 * RATE


QA, INSIGHT = cost(9_794, 2_150), cost(14_341, 4_519)
JUDGE = math.ceil(cost(500_000, 123_000) + 2.6)
HOST, DOMAIN = 1_408, 35
FIXED = HOST + DOMAIN + JUDGE
QA_M, INS_M, FEE = round(100 * QA), round(60 * INSIGHT), round(490 * 0.03)
VARIABLE = QA_M + INS_M + FEE
PRICE = 490
MARGIN = PRICE - VARIABLE
BREAK_EVEN = math.ceil(FIXED / MARGIN)
assert (FIXED, VARIABLE, BREAK_EVEN) == (1456, 53, 4), (FIXED, VARIABLE, BREAK_EVEN)

NAVY, BLUE, ORANGE, GRAY = "14243A", "2073BE", "D46A12", "3F546C"
HEAD_FILL, BAND_FILL, GROUP_FILL, HIGHLIGHT = "E2EEFA", "F3F7FC", "EEF4FB", "FDEEDF"
LINE = "D3DEEA"
FONT = "Microsoft JhengHei"

# ── 移除舊的三張卡片、平衡點方塊與折線圖，保留標題、頁尾與頁碼 ──
KEEP = {"Text 2", "Shape 23", "Text 24", "07"}
for sh in list(s.shapes):
    if sh.name not in KEEP:
        sh._element.getparent().remove(sh._element)


def style_run(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def set_cell(cell, text, size=12, color=NAVY, bold=False, fill="FFFFFF", align=PP_ALIGN.LEFT):
    tf = cell.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = align
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    run = para.add_run()
    run.text = text
    style_run(run, size, color, bold)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Emu(91440)
    cell.margin_top = cell.margin_bottom = Emu(18288)
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(edge), w="9525")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=LINE)
    sf = etree.SubElement(tcPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=fill)


def add_table(rows, cols, left, top, widths, row_h, name):
    gf = s.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(sum(widths)), Inches(row_h * rows))
    gf.name = name
    tbl = gf.table
    # 不套用內建表格樣式（避免首列／帶狀列自動上色），顏色全部由儲存格自行指定。
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.attrib.pop(attr, None)
    style_id = tblPr.find(qn("a:tableStyleId"))
    if style_id is not None:
        tblPr.remove(style_id)
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for r in tbl.rows:
        r.height = Inches(row_h)
    return tbl


def add_text(left, top, width, height, lines, name, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    """lines: [(文字, 字級, 顏色, 粗體, 對齊)]"""
    box = s.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    box.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        box.adjustments[0] = 0.06
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        box.fill.background()
    if line:
        box.line.color.rgb = RGBColor.from_string(line)
        box.line.width = Pt(1)
    else:
        box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    for k, (text, size, color, bold, align) in enumerate(lines):
        para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(3)
        run = para.add_run()
        run.text = text
        style_run(run, size, color, bold)
    return box


# ── 左：成本明細表（項目／計算式／金額） ─────────────────
ROWS = [
    ("head", "項目", "計算式", "元／月"),
    ("group", "固定支出 F：每月都要付，與帳號數無關"),
    ("row", "雲端主機＋資料庫", "（US$25＋US$19）× 匯率 32", f"{HOST:,}"),
    ("row", "網域", "年費 420 ÷ 12 個月", f"{DOMAIN}"),
    ("row", "情緒判讀＋向量", "近 30 天 897 篇＋3,216 則留言", f"{JUDGE}"),
    ("total", "固定支出 F", "以上三項合計", f"{FIXED:,}"),
    ("group", "變動成本 v：每多 1 位專業版用戶，每月增加"),
    ("row", "AI 問答 100 次", f"100 × {QA:.2f} 元", f"{QA_M}"),
    ("row", "儀表板洞察 60 次", f"60 × {INSIGHT:.2f} 元", f"{INS_M}"),
    ("row", "金流手續費", "月費 490 × 3%", f"{FEE}"),
    ("total", "變動成本 v", "以上三項合計（每位用戶／月）", f"{VARIABLE}"),
]
LEFT, TOP, ROW_H = 0.8, 1.78, 0.345
WIDTHS = [2.15, 3.75, 1.2]
cost_tbl = add_table(len(ROWS), 3, LEFT, TOP, WIDTHS, ROW_H, "Cost Table")
band = False
for i, row in enumerate(ROWS):
    kind, *texts = row
    cells = cost_tbl.rows[i].cells
    if kind == "head":
        for j, t in enumerate(texts):
            set_cell(cells[j], t, 13, BLUE, True, HEAD_FILL, PP_ALIGN.RIGHT if j == 2 else PP_ALIGN.LEFT)
    elif kind == "group":
        cells[0].merge(cells[2])
        set_cell(cells[0], texts[0], 12, GRAY, True, GROUP_FILL)
        band = False
    elif kind == "total":
        set_cell(cells[0], texts[0], 13, ORANGE, True, HIGHLIGHT)
        set_cell(cells[1], texts[1], 12, GRAY, False, HIGHLIGHT)
        set_cell(cells[2], texts[2], 14, ORANGE, True, HIGHLIGHT, PP_ALIGN.RIGHT)
    else:
        fill = BAND_FILL if band else "FFFFFF"
        set_cell(cells[0], texts[0], 12, NAVY, False, fill)
        set_cell(cells[1], texts[1], 12, GRAY, False, fill)
        set_cell(cells[2], texts[2], 12, NAVY, True, fill, PP_ALIGN.RIGHT)
        band = not band

# 左下：單次 AI 費用的算法
add_text(LEFT, TOP + ROW_H * len(ROWS) + 0.03, sum(WIDTHS), 0.32, [
    (f"單次費用＝（輸入 token × 0.25＋輸出 token × 1.50）÷ 10⁶ × 32 → 問答 {QA:.2f} 元、洞察 {INSIGHT:.2f} 元",
     10.5, GRAY, False, PP_ALIGN.LEFT),
], "Unit Cost Formula")

# ── 右上：損益公式 ─────────────────────────────────
RX, RW = 8.15, 4.38
add_text(RX, TOP, RW, 1.42, [
    ("損益平衡怎麼算", 14, ORANGE, True, PP_ALIGN.LEFT),
    (f"月收入 R(N)＝{PRICE} × N", 14, NAVY, True, PP_ALIGN.LEFT),
    (f"月支出 C(N)＝{FIXED:,}＋{VARIABLE} × N", 14, NAVY, True, PP_ALIGN.LEFT),
    (f"打平：N ≥ {FIXED:,} ÷（{PRICE}－{VARIABLE}）＝{FIXED / MARGIN:.1f} → {BREAK_EVEN} 個", 13, BLUE, True,
     PP_ALIGN.LEFT),
], "Break-even Formula", fill="E2EDF9", line="C9DAEE", shape=MSO_SHAPE.ROUNDED_RECTANGLE)

# ── 右下：收支對照表（N 個專業版帳號） ────────────────
SCENARIOS = [0, 2, 3, BREAK_EVEN, 5, 10]
# 讓右側表格底緣與左側成本表對齊。
CMP_TOP = TOP + 1.42 + 0.1
CMP_ROW_H = (TOP + ROW_H * len(ROWS) - CMP_TOP) / (len(SCENARIOS) + 1)
cmp_tbl = add_table(len(SCENARIOS) + 1, 4, RX, CMP_TOP, [1.0, 1.08, 1.08, 1.22], CMP_ROW_H, "Break-even Table")
for j, t in enumerate(["帳號數 N", "收入", "支出", "損益"]):
    set_cell(cmp_tbl.rows[0].cells[j], t, 12, BLUE, True, HEAD_FILL, PP_ALIGN.CENTER)
for i, n in enumerate(SCENARIOS, start=1):
    income, spend = n * PRICE, FIXED + n * VARIABLE
    profit = income - spend
    hit = n == BREAK_EVEN
    fill = HIGHLIGHT if hit else ("F3F7FC" if i % 2 == 0 else "FFFFFF")
    cells = cmp_tbl.rows[i].cells
    set_cell(cells[0], f"{n} 個 ✓" if hit else f"{n} 個", 12, ORANGE if hit else NAVY, hit, fill, PP_ALIGN.CENTER)
    set_cell(cells[1], f"{income:,}", 12, NAVY, hit, fill, PP_ALIGN.RIGHT)
    set_cell(cells[2], f"{spend:,}", 12, NAVY, hit, fill, PP_ALIGN.RIGHT)
    sign = "+" if profit > 0 else ("−" if profit < 0 else "")
    color = ORANGE if hit else (BLUE if profit > 0 else GRAY)
    set_cell(cells[3], f"{sign}{abs(profit):,}", 12, color, True, fill, PP_ALIGN.RIGHT)

# ── 頁尾：單價出處 ─────────────────────────────────
for sh in s.shapes:
    if sh.name == "Text 24":
        par = sh.text_frame.paragraphs[0]
        par.runs[0].text = ("新臺幣／月；Gemini 3.1 Flash-Lite 每百萬 token 0.25／1.50 美元（10/18 查，1 USD≈32 元）；"
                            "未含人力，見附錄 C、F。")
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)

# 頁尾與頁碼移到最上層，避免被新加入的圖形蓋住。
tree = s.shapes._spTree
for sh in list(s.shapes):
    if sh.name in {"Shape 23", "Text 24", "07"}:
        tree.remove(sh._element)
        tree.append(sh._element)

p.save(OUT)
print("ok", FIXED, VARIABLE, BREAK_EVEN)
