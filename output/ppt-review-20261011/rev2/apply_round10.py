"""第十輪：統一卡片樣式。

第 2、6、8、13、18 頁各有一張卡片用了橘色粗框（第 8 頁另有橘底），
和同排其他卡片不一致。全部改回一般卡片：白底、1pt 淺灰藍框 D3DEEA；
卡片內原本跟著橘框一起改成橘色的文字，也改回與同排卡片相同的顏色。

輸入 stage9.pptx，輸出 stage10.pptx。
"""
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)

CARD_LINE, CARD_FILL = "D3DEEA", "FFFFFF"
NAVY, BLUE = "14243A", "1F6FB8"

# (頁碼, 卡片 shape_id)
CARDS = [(2, 6), (6, 9), (8, 24), (13, 11), (18, 16)]
# (頁碼, 文字 shape_id, 新顏色)：跟同排卡片一致
TEXT_COLORS = [
    (8, 25, NAVY),   # ST 拉開差異：同 SO／WO／WT 標題
    (8, 27, BLUE),   # ST 驗證列：同其他「驗證：」
    (18, 19, NAVY),  # 面對 AI 代理：同「已完成」「下一步」標題
]


def shape(slide_no, shape_id):
    for sh in p.slides[slide_no - 1].shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError((slide_no, shape_id))


def solid(parent, tag, color):
    el = parent.find(qn(tag))
    if el is None:
        el = etree.SubElement(parent, qn(tag))
    for child in list(el):
        el.remove(child)
    etree.SubElement(el, qn("a:srgbClr"), val=color)
    return el


for slide_no, shape_id in CARDS:
    sp = shape(slide_no, shape_id)
    spPr = sp._element.find(qn("p:spPr"))
    fill = spPr.find(qn("a:solidFill"))
    for child in list(fill):
        fill.remove(child)
    etree.SubElement(fill, qn("a:srgbClr"), val=CARD_FILL)
    ln = spPr.find(qn("a:ln"))
    ln.set("w", "12700")
    solid(ln, "a:solidFill", CARD_LINE)

for slide_no, shape_id, color in TEXT_COLORS:
    sp = shape(slide_no, shape_id)
    for para in sp.text_frame.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor.from_string(color)

p.save(OUT)
print("ok")
