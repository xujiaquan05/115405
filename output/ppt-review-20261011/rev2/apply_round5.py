"""第五輪：刪除第 16 頁（分層測試），並重編頁碼與附錄中的頁碼引用。

輸入 stage4.pptx，輸出 stage5.pptx。
"""
import re
import sys

from pptx import Presentation

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)


def set_text(tf, text):
    par = tf.paragraphs[0]
    par.runs[0].text = text
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)


# 刪除第 16 頁
sld_ids = p.slides._sldIdLst
target = list(sld_ids)[15]
rid = target.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
p.part.drop_rel(rid)
sld_ids.remove(target)

# 每頁左下角的頁碼（兩位數字文字框）依新順序重編
for idx, slide in enumerate(p.slides, start=1):
    for sh in slide.shapes:
        # 只動左下角的頁碼框；第 2 頁議程卡上的 01～04 也是兩位數字，不能一起改
        is_footer = sh.top > 914400 * 6.8 and sh.left < 914400 * 1.0
        if is_footer and sh.has_text_frame and re.fullmatch(r"\d{2}", sh.text_frame.text.strip()):
            set_text(sh.text_frame, f"{idx:02d}")

# 附錄對照表中引用的頁碼：原 17→16、18→17、19→18
shift = {"p17": "p16", "p18": "p17", "p19": "p18"}
for slide in p.slides:
    for sh in slide.shapes:
        if not sh.has_table:
            continue
        for row in sh.table.rows:
            for cell in row.cells:
                for par in cell.text_frame.paragraphs:
                    for run in par.runs:
                        new = re.sub(r"p1[789]", lambda m: shift[m.group(0)], run.text)
                        if new != run.text:
                            run.text = new

# 第 2 頁議程卡：拿掉已刪除的「測試證據」
for sh in p.slides[1].shapes:
    if sh.has_text_frame and "測試證據" in sh.text_frame.text:
        for par in sh.text_frame.paragraphs:
            for run in par.runs:
                run.text = run.text.replace("實機展示、測試證據", "實機展示")

p.save(OUT)
print("slides:", len(p.slides))
