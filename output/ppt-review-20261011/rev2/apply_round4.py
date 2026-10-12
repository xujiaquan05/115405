"""第四輪：投影片上不再顯示報告時間。

- 第 1 頁：「複審 · 15 MIN · 第115405組」→「複審 · 第115405組」
- 第 2 頁：標題「14 分 30 秒主講，保留 30 秒緩衝」改為「報告流程」；四張卡片下方的時段刪除
- 各頁右上角「講者 · 時段」刪除
講者備忘稿中的秒數保留，供排練使用（不會投影出來）。

輸入 stage3.pptx，輸出 stage4.pptx。
"""
import re
import sys

from pptx import Presentation

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)
TIME = re.compile(r"^\s*(\S+ · )?\d+:\d\d–\d+:\d\d\s*$")


def set_text(sh, text):
    par = sh.text_frame.paragraphs[0]
    par.runs[0].text = text
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)


removed = 0
for idx, slide in enumerate(p.slides, start=1):
    for sh in list(slide.shapes):
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text
        if TIME.match(text):
            sh._element.getparent().remove(sh._element)
            removed += 1
        elif idx == 1 and "15 MIN" in text:
            set_text(sh, "複審 · 第115405組")
        elif idx == 2 and "14 分 30 秒" in text:
            set_text(sh, "報告流程")
print("removed time labels:", removed)

p.save(OUT)
print("saved", OUT)
