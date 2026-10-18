"""第八輪：模型改以 Gemini 3.1 Flash-Lite 試算，更新第 7 頁成本與損益平衡。

輸入 stage8-in.pptx（第二輪修訂版），輸出 stage8.pptx。

單價：gemini-3.1-flash-lite 輸入 0.25／輸出 1.50 美元／百萬 token（ai.google.dev，10/18 查）。
token 數沿用 backend/scripts/measure_llm_cost.py 的實測值，匯率 32。
"""
import math
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)

RATE = 32
IN_PRICE, OUT_PRICE = 0.25, 1.50  # 美元／百萬 token


def cost(tokens_in, tokens_out):
    return (tokens_in * IN_PRICE + tokens_out * OUT_PRICE) / 1e6 * RATE


QA = cost(9_794, 2_150)              # 每次問答
INSIGHT = cost(14_341, 4_519)        # 每次洞察
SENTIMENT = cost(500_000, 123_000)   # 每月情緒判讀
EMBEDDING = 2.6                      # 向量（gemini-embedding-001，模型不變）
JUDGE = math.ceil(SENTIMENT + EMBEDDING)
FIXED = 1_408 + 35 + JUDGE
VARIABLE = round(100 * QA + 60 * INSIGHT + 490 * 0.03)
PRICE = 490
BREAK_EVEN = math.ceil(FIXED / (PRICE - VARIABLE))
INCOME = BREAK_EVEN * PRICE
SPEND = FIXED + BREAK_EVEN * VARIABLE
LABOR = [math.ceil((FIXED + hours * 196) / (PRICE - VARIABLE)) for hours in (20, 40)]

print(f"QA={QA:.3f} insight={INSIGHT:.3f} sentiment={SENTIMENT:.1f} fixed={FIXED} "
      f"v={VARIABLE} N={BREAK_EVEN} income={INCOME} spend={SPEND} labor={LABOR}")


def set_run(sh, text):
    par = sh.text_frame.paragraphs[0]
    par.runs[0].text = text
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)


s = p.slides[6]
replacements = {
    "成本重算：4 個專業版帳號可打平雲端支出": f"成本重算：{BREAK_EVEN} 個專業版帳號可打平雲端支出",
    "0.06 元": f"{QA:.2f} 元",
    "27 元": f"{VARIABLE} 元",
    "1,450 元": f"{FIXED:,} 元",
    "4 個": f"{BREAK_EVEN} 個",
    "收入 1,960 ＞ 支出 1,558": f"收入 {INCOME:,} ＞ 支出 {SPEND:,}",
    "每月 1,450 固定支出＋每帳號 27": f"每月 {FIXED:,} 固定支出＋每帳號 {VARIABLE}",
    "新臺幣／月；官方單價查於 10/12（1 USD≈32 元）；未含人力，見附錄 C、F。":
        "新臺幣／月；Gemini 3.1 Flash-Lite 官方單價查於 10/18（1 USD≈32 元）；未含人力，見附錄 C、F。",
}
done = set()
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text in replacements:
        set_run(sh, replacements[sh.text_frame.text])
        done.add(sh.text_frame.text)
missing = [k for k in replacements if replacements[k] not in done and k not in done]
assert len(done) == len(replacements), missing

chart = next(sh for sh in s.shapes if sh.has_chart).chart
accounts = [0, 2, 4, 6, 8, 10]
cd = CategoryChartData()
cd.categories = [str(n) for n in accounts]
cd.add_series("月收入", [n * PRICE for n in accounts])
cd.add_series("月支出", [FIXED + n * VARIABLE for n in accounts])
chart.replace_data(cd)

s.notes_slide.notes_text_frame.paragraphs[0].runs[0].text = (
    "【陳威帆｜45 秒】初評說我們沒有量化，這頁用實測重算，不再沿用原稿的假設單價。"
    "一次 AI 問答包含問題解析與依證據回答兩次呼叫，用系統程式組出真實 prompt，最多約 9,800 token 輸入、2,150 token 輸出，"
    f"以 Gemini 3.1 Flash-Lite 官方價每百萬 0.25／1.50 美元計算，約 {QA:.2f} 元；一次洞察約 {INSIGHT:.2f} 元。"
    f"專業版用戶每月 100 次問答、60 次洞察，加上 3% 金流手續費，變動成本約 {VARIABLE} 元。"
    f"固定支出以 render.yaml 的雲端部署計：主機 25 美元、資料庫 19 美元，加網域與每日判讀約 {FIXED:,} 元。"
    f"收入 490N、支出 {FIXED:,}＋{VARIABLE}N，{BREAK_EVEN} 個帳號就能打平，每月餘 {INCOME - SPEND} 元。"
    f"但這還沒算人力，算進去要 {LABOR[0]} 到 {LABOR[1]} 個帳號，見附錄 C；每個數字的出處在附錄 F。"
)

p.save(OUT)
