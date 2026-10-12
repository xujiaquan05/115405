"""第七輪：AI 問答頁改為真正走混合檢索後，同步簡報的數字與備忘稿。

輸入 stage6.pptx，輸出 stage7.pptx。
"""
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)


def note(idx, text):
    p.slides[idx - 1].notes_slide.notes_text_frame.paragraphs[0].runs[0].text = text


def set_run(sh, text):
    par = sh.text_frame.paragraphs[0]
    par.runs[0].text = text
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)


# 第 7 頁備忘稿：新流程含問題解析與 Dashboard 摘要
note(7, "【陳威帆｜45 秒】初評說我們沒有量化，這頁用實測重算，不再沿用原稿的假設單價。"
        "一次 AI 問答包含問題解析與依證據回答兩次呼叫，用系統程式組出真實 prompt，最多約 9,800 token 輸入、2,150 token 輸出，"
        "以 Gemini 2.5 Flash-Lite 官方價每百萬 0.10／0.40 美元計算，約 0.06 元，原稿假設 0.33 元高估了五倍。"
        "專業版用戶每月 100 次問答、60 次洞察，加上 3% 金流手續費，變動成本約 27 元。"
        "固定支出以 render.yaml 的雲端部署計：主機 25 美元、資料庫 19 美元，加網域與每日判讀約 1,450 元。"
        "收入 490N、支出 1,450＋27N，4 個帳號就能打平，每月餘 402 元。但這還沒算人力，算進去要 12 到 21 個帳號，見附錄 C；每個數字的出處在附錄 F。")

# 第 11 頁備忘稿：說明問答頁如何使用 RAG
note(11, "【許家全｜45 秒】初評指出只有 keyword 比對。現在先解析問題意圖，再以關鍵字和向量雙路找證據，RRF 合併排序後給模型作答。"
         "在 AI 問答頁提問時，檢索範圍會套用使用者在 Dashboard 選的關鍵字與期間，回答以檢索到的文章為證據、Dashboard 數字為背景，引用清單就是這些文章。"
         "資料預先切段並建立向量；向量無結果時回退關鍵字，完全找不到文章時才只用 Dashboard 數字回答。"
         "引用支持人工回查，不保證每句結論都有充分證據；尚未完成專家知識庫和固定題組評估。"
         "來源：backend/app/services/rag_service.py（answer_question、scope_intent_to_dashboard、RRF_K）；embedding_service.py。")

# 第 16 頁：10 月里程碑與提交數（含本輪提交）
s = p.slides[15]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "情緒一致性、預警權限修正":
        set_run(sh, "情緒一致性、預警與問答修正")
MONTHS = [14, 0, 27, 42, 47, 36]
chart = next(sh for sh in s.shapes if sh.has_chart).chart
cd = CategoryChartData()
cd.categories = ["5月", "6月", "7月", "8月", "9月", "10月"]
cd.add_series("提交數", MONTHS)
chart.replace_data(cd)
chart.chart_title.text_frame.paragraphs[0].runs[0].text = f"GitHub 每月提交（共 {sum(MONTHS)} 次）"
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("GitHub："):
        set_run(sh, f"GitHub：7–10 月 {sum(MONTHS[2:])} 次提交（10/12 統計，皆已推送）；非逐人貢獻比例。")
note(16, "【陳威帆｜45 秒】甘特圖保留原稿115/2至116/1的計畫，10月後為計畫工作。繳件與展覽日期沿用原稿，發表前仍按學校公告確認。"
         "右側里程碑依各功能第一次提交的月份：5月PTT爬蟲、儀表板與WebSocket；7月帳號登入；8月Dcard、Mobile01、Threads；"
         "9月方案額度、Alembic與向量檢索；10月情緒一致性、預警權限與問答檢索修正。"
         f"提交數由GitHub於10/12統計：5月14、6月0、7月27、8月42、9月47、10月36，共{sum(MONTHS)}次，皆已推送；"
         "提交集中於單一帳號，不能當作兩位組員的貢獻比例。")

# 第 18 頁備忘稿：測試數
n = p.slides[17].notes_slide.notes_text_frame.paragraphs[0].runs[0]
n.text = n.text.replace("後端493項、前端43項自動化測試通過", "後端498項、前端43項自動化測試通過")

# 附錄 F：問答成本列
t = next(sh for sh in p.slides[23].shapes if sh.has_table).table
cell = t.cell(1, 1).text_frame.paragraphs[0]
cell.runs[0].text = "兩次呼叫：輸入 9,794＋輸出 2,150 ≈ 0.06 元"
for r in cell.runs[1:]:
    r._r.getparent().remove(r._r)
cell = t.cell(1, 0).text_frame.paragraphs[0]
cell.runs[0].text = "AI 問答／次"
for r in cell.runs[1:]:
    r._r.getparent().remove(r._r)

p.save(OUT)
print("saved", OUT, sum(MONTHS))
