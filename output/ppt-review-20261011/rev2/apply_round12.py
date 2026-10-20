"""第十二輪：第 16–18 頁（第四部分「怎麼管、下一步」）改由許家全主講，備忘稿換成完整講稿。

輸入使用者 10/20 手動修改後的正式版，輸出 stage12.pptx；只改這三頁的備忘稿，其餘內容不動。
"""
import copy
import sys

from pptx import Presentation

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)

NOTES = {
    16: [
        "【許家全｜45 秒】",
        "接下來由我說明進度管理與後續規劃。",
        "初評時老師提醒甘特圖要用實際年月，所以這頁從 115 年 2 月排到 116 年 1 月，分成起始、精化、建構、移轉四個階段，"
        "目前正在移轉階段。上方標出四個查核點：6 月 2 日初審、10 月 13 日複審繳件、11 月 25 日觀摩展，以及 12 月 29 日最終版。",
        "左邊是 GitHub 每月提交數，統計到 10 月 12 日共 166 次；右邊是各月份第一次完成的功能：5 月 PTT 爬蟲與儀表板，"
        "7 月帳號與權限，8 月擴充到 Dcard、Mobile01、Threads，9 月加入方案額度、資料庫遷移與向量檢索，"
        "10 月則完成情緒一致性與問答修正。",
        "來源：GitHub 提交紀錄（10/12 統計：5 月 14、6 月 0、7 月 27、8 月 42、9 月 47、10 月 36）。",
    ],
    17: [
        "【許家全｜30 秒】",
        "分工上，組長威帆負責需求企劃與雜誌社交流，以及成本、STP、TOWS 等營運計畫，並統整系統手冊第 1 到 6 章與簡報。",
        "我負責系統架構、資料庫與 FastAPI 後端，四平台爬蟲和每日排程，AI 問答、情緒評分與向量檢索，"
        "也包含前端介面、測試，以及手冊第 7 到 12 章。",
        "另外說明，程式提交集中在單一帳號，所以提交次數不能直接當作兩個人的貢獻比例。",
        "來源：使用者原稿第 16 頁工作分配；本機 Git 提交歷史。",
    ],
    18: [
        "【許家全｜105 秒】",
        "最後做個總結。",
        "先說已經完成的：系統每天自動收錄四個平台的討論；AI 問答用關鍵字加向量的混合檢索，每個回答都附引用來源；"
        "情緒評分有 100 篇重跑的一致性紀錄；帳號、方案額度與監測預警都在後端檢查，目前後端 528 項、前端 47 項自動化測試全部通過。"
        "營運面依實測重算，4 個專業版帳號就能打平每月約 1,456 元的雲端支出；如果算入每月 20 到 40 小時的維運人力，"
        "則需要 13 到 22 個帳號。",
        "但我們也要清楚說明還差什麼。下一步依序是：補上 Linux 部署的回歸驗證；建立人工標註，驗證情緒判讀的準確度；"
        "用固定題組評估檢索品質；開發業配辨識；最後和實際使用者進行正式的需求訪談與試用。",
        "至於評審提到的通用 AI 代理：它確實能自己搜尋、整理資料，我們不宣稱一定比它好。MeBOD 的定位是長期累積的歷史資料、"
        "每個回答都能回查原文，以及依帳號持續監測與預警；兩者的差異，我們會用同一題實際比較來驗證。",
        "以上是我們的報告，謝謝老師，敬請指教。",
        "備註：測試數為 10/20 統計（後端 528 項通過、1 項略過；前端 47 項）。",
    ],
}


def set_notes(slide_no, paragraphs):
    """以第一段的格式為範本，把備忘稿換成多段文字。"""
    tf = p.slides[slide_no - 1].notes_slide.notes_text_frame
    first = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for r in first.runs[1:]:
        r._r.getparent().remove(r._r)
    template = copy.deepcopy(first._p)

    first.runs[0].text = paragraphs[0]
    anchor = first._p
    for text in paragraphs[1:]:
        new_p = copy.deepcopy(template)
        anchor.addnext(new_p)
        anchor = new_p
    for para, text in zip(tf.paragraphs, paragraphs):
        para.runs[0].text = text
    assert [para.text for para in tf.paragraphs] == paragraphs, slide_no


for slide_no, paragraphs in NOTES.items():
    set_notes(slide_no, paragraphs)

p.save(OUT)
print("ok")
