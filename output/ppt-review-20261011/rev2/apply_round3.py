"""第三輪：封面加入海報 1 的主視覺插畫，並刪除各頁左上角的橘色小標。

輸入 stage2.pptx，輸出 stage3.pptx。
"""
import sys

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def add_shadow(pic, blur, dist, alpha):
    spPr = pic._element.spPr
    eff = spPr.makeelement(A + "effectLst", {})
    shdw = eff.makeelement(A + "outerShdw", {"blurRad": str(blur), "dist": str(dist), "dir": "5400000",
                                             "algn": "t", "rotWithShape": "0"})
    clr = shdw.makeelement(A + "srgbClr", {"val": "4F46E5"})
    clr.append(clr.makeelement(A + "alpha", {"val": str(alpha)}))
    shdw.append(clr)
    eff.append(shdw)
    spPr.append(eff)


# ── 封面 ──
s = p.slides[0]
for sh in list(s.shapes):
    if sh.name == "MeBOD Logo":
        sh._element.getparent().remove(sh._element)

# 插畫放在右側、滿版高度；左緣已預先做成透明漸層，與底色銜接
hero = s.shapes.add_picture("cover-hero.png", Inches(6.9), 0, height=p.slide_height)
hero.name = "Cover Illustration"
tree = s.shapes._spTree
tree.remove(hero._element)
tree.insert(2, hero._element)  # 排在所有文字之後繪製之前，當作背景

# 標誌縮小、放在「MeBOD」左邊，與海報 1 的排法一致
title = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text == "MeBOD")
logo = s.shapes.add_picture("mebod-mark.png", Inches(0.8), title.top + Inches(0.12), Inches(0.95), Inches(0.95))
logo.name = "MeBOD Logo"
add_shadow(logo, 152400, 50800, 30000)
title.left = Inches(1.95)
title.width = Inches(4.9)

# 文字框不要延伸到插畫底下
for sh in s.shapes:
    if sh.has_text_frame and sh.left < Inches(1) and sh.width > Inches(6.1):
        sh.width = Inches(6.1)

# 組員一行放不下（右側是插畫），拆成「組員」與「指導老師」兩行
import copy
team = next(sh for sh in s.shapes if sh.has_text_frame and "指導老師" in sh.text_frame.text)
par = team.text_frame.paragraphs[0]
for r in par.runs[1:]:
    r._r.getparent().remove(r._r)
par.runs[0].text = "組長 陳威帆 11246044　｜　組員 許家全 11246039"
second = copy.deepcopy(par._p)
par._p.addnext(second)
team.text_frame.paragraphs[1].runs[0].text = "指導老師 陳信宏 老師"
team.top = Inches(5.75)
team.height = Inches(0.8)

# ── 刪除左上角橘色小標（第 2 頁起；封面的「複審 · 15 MIN · 第115405組」保留） ──
removed = 0
for idx, slide in enumerate(p.slides, start=1):
    if idx == 1:
        continue
    for sh in list(slide.shapes):
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        near_corner = abs(sh.left - Inches(0.8)) < Inches(0.05) and sh.top < Inches(0.6) and sh.height < Inches(0.4)
        if near_corner:
            sh._element.getparent().remove(sh._element)
            removed += 1
print("removed kickers:", removed)

# 第 2 頁的說明原本指向左上角小標，改成指向附錄對照表
s2 = p.slides[1]
for sh in s2.shapes:
    if sh.has_text_frame and "左上角" in sh.text_frame.text:
        run = sh.text_frame.paragraphs[0].runs[0]
        run.text = "16 項初評意見與回應頁碼，逐條對照於附錄 A、B"
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r._r.getparent().remove(r._r)

# 第 7 頁損益平衡框：直接寫出「4 個什麼」以及收支比較，不必聽講解也看得懂
s7 = p.slides[6]
def put(name, text):
    sh = next(x for x in s7.shapes if x.name == name)
    par = sh.text_frame.paragraphs[0]
    par.runs[0].text = text
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)
    return sh
put("Text 17", "損益平衡點")
put("Text 18", "4 個")
put("Text 19", "專業版付費帳號／月")
put("Text 20", "收入 1,960 ＞ 支出 1,558")
put("Text 21", "每月 1,450 固定支出＋每帳號 27")
for n in ("Text 19", "Text 20"):
    next(x for x in s7.shapes if x.name == n).text_frame.paragraphs[0].runs[0].font.size = Pt(14)
next(x for x in s7.shapes if x.name == "Text 19").text_frame.paragraphs[0].runs[0].font.bold = True

# 小標刪除後，學校規定的段落名稱與附錄代號改寫進標題，避免「見附錄 C、F」找不到頁
PREFIX = {
    4: ("初評意見已轉成具體修正與待辦", "前次評審意見修正情形：初評意見已轉成具體修正與待辦"),
    14: ("一條研究流程，對應六組可操作功能", "功能完整性：一條研究流程，對應六組可操作功能"),
    20: ("16 項意見對照：定位、營運與實作（1/2）", "附錄 A｜16 項意見對照：定位、營運與實作（1/2）"),
    21: ("16 項意見對照：資料品質與驗證（2/2）", "附錄 B｜16 項意見對照：資料品質與驗證（2/2）"),
    22: ("加入維運人力後，平衡帳號數變成 12～21 個", "附錄 C｜加入維運人力後，平衡帳號數變成 12～21 個"),
    23: ("AI 問答：從回答回到引用文章", "附錄 D｜AI 問答：從回答回到引用文章"),
    24: ("歷史比較：保留兩次查詢的分析結果", "附錄 E｜歷史比較：保留兩次查詢的分析結果"),
    25: ("每一個數字的算法與出處", "附錄 F｜每一個數字的算法與出處"),
}
for idx, (old, new) in PREFIX.items():
    sh = next(x for x in p.slides[idx - 1].shapes if x.has_text_frame and x.text_frame.text.strip() == old)
    par = sh.text_frame.paragraphs[0]
    par.runs[0].text = new
    for r in par.runs[1:]:
        r._r.getparent().remove(r._r)

p.save(OUT)
print("saved", OUT)
